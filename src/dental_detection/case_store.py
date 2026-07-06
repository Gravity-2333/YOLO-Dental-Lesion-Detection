from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from .assistant import SAFETY_NOTICE, case_dir, ensure_app_dirs, report_dir, storage_root
from .model_info import legend_markdown
from .result_levels import REGION_NOTICE, enrich_detection_row, get_class_display_name, parse_confidence
from .text_utils import text_value


def _case_trash_dir(storage_dir: str) -> Path:
    return storage_root(storage_dir) / "cases_trash"


def _safe_case_file_name(file_name: str) -> str:
    name = Path(str(file_name or "")).name
    if name != file_name or not name.startswith("case_") or not name.endswith(".json"):
        raise ValueError("病例文件名无效。")
    return name


def _read_case_file(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def _case_detections(data: dict[str, Any]) -> list[dict[str, Any]]:
    model_results = data.get("model_results")
    if isinstance(model_results, list) and model_results:
        rows = []
        for result in model_results:
            if not isinstance(result, dict):
                continue
            model_name = result.get("model") or result.get("模型")
            for det in result.get("detections") or []:
                if not isinstance(det, dict):
                    continue
                row = enrich_detection_row(det)
                if model_name:
                    row["模型"] = str(model_name)
                rows.append(row)
        if rows:
            return rows
    return [enrich_detection_row(det) for det in data.get("detections") or [] if isinstance(det, dict)]


def _case_classes(detections: list[dict[str, Any]]) -> list[str]:
    classes = []
    for det in detections:
        display = det.get("中文名称") or get_class_display_name(det.get("class", ""))
        if display not in classes:
            classes.append(display)
    return classes


def _case_levels(detections: list[dict[str, Any]]) -> list[str]:
    levels = []
    for det in detections:
        level = str(det.get("关注等级") or "")
        if level and level not in levels:
            levels.append(level)
    return levels


def _highest_confidence(detections: list[dict[str, Any]]) -> Any:
    values = [parse_confidence(det.get("confidence")) for det in detections]
    values = [value for value in values if value is not None]
    return round(max(values), 4) if values else ""


def _row_from_case(path: Path, data: dict[str, Any]) -> dict[str, Any]:
    detections = _case_detections(data)
    classes = _case_classes(detections)
    levels = _case_levels(detections)
    return {
        "保存时间": text_value(data.get("created_at")),
        "病例编号": text_value(data.get("case_id"), "未填写"),
        "图片名称": text_value(data.get("display_name") or data.get("image_name")),
        "检测数量": len(detections),
        "涉及类别": "、".join(classes) if classes else "无检测结果",
        "关注等级": "、".join(levels) if levels else "无检测结果",
        "最高置信度": _highest_confidence(detections),
        "文件名": path.name,
        "_data": data,
    }


def list_case_records(storage_dir: str) -> list[dict[str, Any]]:
    ensure_app_dirs(storage_dir)
    rows: list[dict[str, Any]] = []
    for path in sorted(case_dir(storage_dir).glob("case_*.json"), reverse=True):
        try:
            rows.append(_row_from_case(path, _read_case_file(path)))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            rows.append(
                {
                    "保存时间": "",
                    "病例编号": "损坏病例文件",
                    "图片名称": "",
                    "检测数量": "",
                    "涉及类别": "",
                    "关注等级": "",
                    "最高置信度": "",
                    "文件名": path.name,
                    "_data": {"错误": f"病例文件损坏或无法读取：{path.name}"},
                }
            )
    return rows


def search_case_records(
    storage_dir: str,
    keyword: str,
    class_filter: str,
    level_filter: str,
    date_from: str,
    date_to: str,
) -> list[dict[str, Any]]:
    keyword_text = str(keyword or "").strip().casefold()
    class_value = str(class_filter or "全部")
    level_value = str(level_filter or "全部")
    start = str(date_from or "").strip()
    end = str(date_to or "").strip()
    rows = []
    for row in list_case_records(storage_dir):
        data = row.get("_data", {})
        detections = _case_detections(data)
        searchable = "\n".join(
            str(value)
            for value in [
                row.get("病例编号", ""),
                row.get("图片名称", ""),
                data.get("note", ""),
                data.get("suggestion", ""),
                row.get("涉及类别", ""),
                row.get("关注等级", ""),
            ]
        ).casefold()
        if keyword_text and keyword_text not in searchable:
            continue
        if class_value != "全部":
            if class_value == "无检测结果":
                if detections:
                    continue
            elif class_value not in row.get("涉及类别", ""):
                continue
        if level_value != "全部":
            if level_value == "无检测结果":
                if detections:
                    continue
            elif level_value not in row.get("关注等级", ""):
                continue
        created_at = str(row.get("保存时间") or "")
        if start and created_at[:10] < start:
            continue
        if end and created_at[:10] > end:
            continue
        rows.append(row)
    return rows


def load_case_record(storage_dir: str, file_name: str) -> dict[str, Any]:
    safe_name = _safe_case_file_name(file_name)
    root = case_dir(storage_dir).resolve()
    path = (root / safe_name).resolve()
    if path.parent != root:
        raise ValueError("病例文件路径无效。")
    if not path.exists():
        raise FileNotFoundError(f"病例文件不存在：{safe_name}")
    return _read_case_file(path)


def move_case_to_trash(storage_dir: str, file_name: str) -> Path:
    safe_name = _safe_case_file_name(file_name)
    ensure_app_dirs(storage_dir)
    source_root = case_dir(storage_dir).resolve()
    source = (source_root / safe_name).resolve()
    if source.parent != source_root:
        raise ValueError("病例文件路径无效。")
    if not source.exists():
        raise FileNotFoundError(f"病例文件不存在：{safe_name}")
    trash = _case_trash_dir(storage_dir)
    trash.mkdir(parents=True, exist_ok=True)
    target = trash / safe_name
    if target.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        target = trash / f"{Path(safe_name).stem}_{stamp}.json"
    shutil.move(str(source), str(target))
    return target


def export_case_report(storage_dir: str, file_name: str) -> Path:
    data = load_case_record(storage_dir, file_name)
    output_dir = report_dir(storage_dir) / "case_reports"
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"病例记录报告_{stamp}_{Path(file_name).stem}.docx"
    counter = 1
    while output_path.exists():
        output_path = output_dir / f"病例记录报告_{stamp}_{Path(file_name).stem}_{counter:02d}.docx"
        counter += 1

    from docx import Document

    detections = _case_detections(data)
    document = Document()
    document.add_heading("牙齿病变区域辅助识别病例记录", level=0)
    document.add_paragraph(SAFETY_NOTICE)
    document.add_paragraph("该病例记录仅保存检测摘要、检测框和辅助建议，不包含原始牙片图片。")
    document.add_paragraph(f"保存时间：{data.get('created_at', '-')}")
    document.add_paragraph(f"病例编号：{data.get('case_id', '未填写')}")
    image_name = text_value(data.get("image_name"), "-")
    display_name = text_value(data.get("display_name")).strip()
    document.add_paragraph(f"图片名称：{image_name}")
    if display_name and display_name != image_name:
        document.add_paragraph(f"列表显示名：{display_name}")
    if data.get("note"):
        document.add_paragraph(f"病例备注：{data.get('note')}")
    report_path = data.get("report_path") or data.get("word_report_path") or data.get("zip_report_path")
    if report_path:
        document.add_paragraph(f"关联报告路径：{report_path}")

    document.add_heading("检测框明细", level=1)
    document.add_paragraph(REGION_NOTICE)
    headers = ["序号", "模型", "类别", "中文名称", "置信度", "关注等级", "图像区域", "x1", "y1", "x2", "y2"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    if detections:
        for index, row in enumerate(detections, start=1):
            values = [
                index,
                row.get("模型", ""),
                row.get("class", ""),
                row.get("中文名称", ""),
                row.get("confidence", ""),
                row.get("关注等级", ""),
                row.get("图像区域", ""),
                row.get("x1", ""),
                row.get("y1", ""),
                row.get("x2", ""),
                row.get("y2", ""),
            ]
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
    else:
        for cell, value in zip(table.add_row().cells, ["-", "", "无检测结果", "", "", "", "", "", "", "", ""]):
            cell.text = str(value)

    if data.get("quality_text"):
        document.add_heading("图像质量提示", level=1)
        for line in str(data.get("quality_text")).splitlines():
            document.add_paragraph(line)

    if data.get("suggestion"):
        document.add_heading("辅助建议", level=1)
        for line in str(data.get("suggestion")).splitlines():
            document.add_paragraph(line)

    document.add_heading("图例与类别说明", level=1)
    for line in legend_markdown().splitlines():
        document.add_paragraph(line)

    document.save(output_path)
    return output_path
