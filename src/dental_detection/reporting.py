from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from PIL import Image

from .batch_summary import build_batch_summary
from .model_info import legend_markdown
from .result_levels import REGION_NOTICE, enrich_detection_row, has_detection_payload, parse_confidence
from .result_items import item_model_results
from .visualization import save_png_image


@dataclass
class SingleReportData:
    image_name: str
    created_at: str
    model_name: str
    original_image: Image.Image
    model_input_image: Image.Image | None
    annotated_image: Image.Image
    detections: list[dict[str, Any]]
    advice: str
    quality_text: str
    summary: dict[str, Any]
    safety_notice: str
    model_results: list[dict[str, Any]] = field(default_factory=list)


def _add_paragraphs(document, text: str) -> None:
    for line in str(text or "").splitlines() or [""]:
        document.add_paragraph(line)


def _save_temp_image(image: Image.Image, root: Path, name: str) -> Path:
    path = root / name
    return save_png_image(image, path)


def _summary_value(summary: dict[str, Any], *names: str, default: str = "-") -> Any:
    for name in names:
        if name in summary:
            return summary[name]
    return default


def _yes_no(value: Any) -> str:
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _image_size(image: Any) -> tuple[int, int] | None:
    size = getattr(image, "size", None)
    if isinstance(size, tuple) and len(size) == 2:
        try:
            return int(size[0]), int(size[1])
        except (TypeError, ValueError):
            return None
    shape = getattr(image, "shape", None)
    if shape is not None and len(shape) >= 2:
        try:
            return int(shape[1]), int(shape[0])
        except (TypeError, ValueError):
            return None
    return None


def export_single_docx_report(data: SingleReportData, output_dir: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "牙齿病变区域辅助识别报告.docx"
    counter = 1
    while output_path.exists():
        output_path = output_dir / f"牙齿病变区域辅助识别报告_{counter:02d}.docx"
        counter += 1

    document = Document()
    document.add_heading("牙齿病变区域辅助识别报告", level=0)
    notice = document.add_paragraph()
    run = notice.add_run(data.safety_notice)
    run.bold = True

    document.add_paragraph(f"生成时间：{data.created_at}")
    document.add_paragraph(f"图片名称：{data.image_name}")
    model_groups = _single_report_model_results(data)
    model_names = [item["model"] for item in model_groups if item.get("model")]
    document.add_paragraph(f"使用模型：{'、'.join(dict.fromkeys(model_names)) if model_names else data.model_name}")
    document.add_paragraph("报告类型：单张影像检测报告")

    document.add_heading("一、影像与检测结果", level=1)
    with TemporaryDirectory() as temp_dir:
        temp_root = Path(temp_dir)
        for caption, image, file_name in [
            ("图 1 原始上传图", data.original_image, "original.png"),
            ("图 2 检测结果图", data.annotated_image, "annotated.png"),
            ("图 3 实际送入模型的图像", data.model_input_image, "model_input.png"),
        ]:
            if image is None:
                continue
            document.add_paragraph(caption)
            image_path = _save_temp_image(image, temp_root, file_name)
            document.add_picture(str(image_path), width=Inches(5.8))
        if len(model_groups) > 1:
            for index, group in enumerate(model_groups, start=1):
                image = group.get("annotated_image")
                if image is None:
                    continue
                model_name = group.get("model") or f"模型 {index}"
                document.add_paragraph(f"图 {index + 3} 分模型结果图：{model_name}")
                image_path = _save_temp_image(image, temp_root, f"model_{index:02d}.png")
                document.add_picture(str(image_path), width=Inches(5.8))

    document.add_heading("二、检测摘要", level=1)
    detection_image_size = _image_size(data.model_input_image) or _image_size(data.original_image)
    enriched = []
    for group in model_groups:
        model_name = group.get("model", "")
        for det in group.get("detections", []):
            if not has_detection_payload(det):
                continue
            row = enrich_detection_row(det, detection_image_size)
            row["模型"] = model_name
            enriched.append(row)
    if enriched:
        document.add_paragraph(f"本次共检测到 {len(enriched)} 个疑似目标区域。")
        if len(model_groups) > 1:
            document.add_paragraph(f"模型结果组数：{len(model_groups)}")
        counts: dict[str, int] = {}
        for row in enriched:
            display_name = row["中文名称"]
            counts[display_name] = counts.get(display_name, 0) + 1
        for name, count in sorted(counts.items()):
            document.add_paragraph(f"- {name}：{count} 处")
    else:
        document.add_paragraph(
            "本次未检测到明确目标框。该结果不代表不存在病变，仍需结合原始影像质量、临床症状和专业牙科检查进行判断。"
        )

    document.add_heading("三、检测框明细", level=1)
    document.add_paragraph(REGION_NOTICE)
    headers = ["序号", "模型", "类别", "中文名称", "置信度", "关注等级", "图像区域", "x1", "y1", "x2", "y2"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    if enriched:
        for index, row in enumerate(enriched, start=1):
            cells = table.add_row().cells
            values = [
                index,
                row.get("模型", ""),
                row["class"],
                row["中文名称"],
                row["confidence"],
                row["关注等级"],
                row.get("图像区域", ""),
                row["x1"],
                row["y1"],
                row["x2"],
                row["y2"],
            ]
            for cell, value in zip(cells, values):
                cell.text = str(value)
    else:
        for cell, value in zip(table.add_row().cells, ["-", "", "未检测到目标框", "", "", "", "", "", "", "", ""]):
            cell.text = str(value)

    document.add_heading("四、图像质量提示", level=1)
    _add_paragraphs(document, data.quality_text)

    document.add_heading("五、辅助建议", level=1)
    _add_paragraphs(document, data.advice)

    document.add_heading("六、图例与类别说明", level=1)
    _add_paragraphs(document, legend_markdown())

    document.add_heading("七、技术参数摘要", level=1)
    document.add_paragraph(f"推理尺寸：{_summary_value(data.summary, '推理尺寸')}")
    document.add_paragraph(f"置信度阈值：{_summary_value(data.summary, '置信度阈值', 'conf')}")
    document.add_paragraph(f"IoU 阈值：{_summary_value(data.summary, 'IoU阈值', 'iou')}")
    document.add_paragraph(f"CLAHE 增强：{_yes_no(_summary_value(data.summary, 'CLAHE增强'))}")
    document.add_paragraph(f"运行设备：{_summary_value(data.summary, '运行设备')}")
    document.add_paragraph(f"模型名称：{'、'.join(dict.fromkeys(model_names)) if model_names else data.model_name}")

    document.save(output_path)
    return output_path


def _safe_name(name: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "._- " else "_" for ch in str(name or "image")).strip() or "image"


def _item_results(item: dict[str, Any]) -> list[dict[str, Any]]:
    return item_model_results(item)


def _primary_result(item: dict[str, Any]) -> dict[str, Any]:
    result = item.get("result") or item
    return result if isinstance(result, dict) else {}


def _first_present(*values: Any) -> Any:
    for value in values:
        if value is not None:
            return value
    return None


def _model_result_image(item: dict[str, Any]) -> Any:
    return _first_present(item.get("full_annotated"), item.get("annotated"), item.get("result_image"))


def _model_name(item: dict[str, Any]) -> str:
    result = _primary_result(item)
    if result.get("model"):
        return str(result["model"])
    summary = item.get("summary")
    if isinstance(summary, dict):
        model_results = summary.get("模型结果")
        if isinstance(model_results, list) and model_results:
            first = model_results[0]
            if isinstance(first, dict) and first.get("模型"):
                return str(first["模型"])
        return str(summary.get("模型", "unknown"))
    return "unknown"


def _single_report_model_results(data: SingleReportData) -> list[dict[str, Any]]:
    if data.model_results:
        rows = []
        for index, item in enumerate(data.model_results, start=1):
            if not isinstance(item, dict):
                continue
            rows.append(
                {
                    "model": str(item.get("model") or item.get("模型") or f"模型 {index}"),
                    "model_path": str(item.get("model_path") or item.get("路径") or item.get("模型路径") or ""),
                    "detections": [det for det in item.get("detections", []) if has_detection_payload(det)],
                    "annotated_image": _model_result_image(item),
                }
            )
        if rows and (any(item.get("detections") for item in rows) or not data.detections):
            return rows
    return [{"model": data.model_name, "model_path": "", "detections": data.detections, "annotated_image": data.annotated_image}]


def _add_detection_table(document, detections: list[dict[str, Any]], image_size: tuple[int, int] | None = None) -> None:
    document.add_paragraph(REGION_NOTICE)
    headers = ["序号", "类别", "中文名称", "置信度", "关注等级", "图像区域", "x1", "y1", "x2", "y2"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    rows = [enrich_detection_row(det, image_size) for det in detections if has_detection_payload(det)]
    if not rows:
        cells = table.add_row().cells
        cells[0].text = "-"
        cells[1].text = "未检测到目标框"
        return
    for index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        values = [
            index,
            row["class"],
            row["中文名称"],
            row["confidence"],
            row["关注等级"],
            row.get("图像区域", ""),
            row["x1"],
            row["y1"],
            row["x2"],
            row["y2"],
        ]
        for cell, value in zip(cells, values):
            cell.text = str(value)


def export_batch_docx_report(
    batch_state: list[dict[str, Any]],
    batch_summary: dict[str, Any] | None,
    output_dir: Path,
) -> Path:
    from docx import Document
    from docx.shared import Inches

    if not batch_state:
        raise ValueError("批量结果为空，无法导出合并报告。")

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "牙齿病变批量辅助识别报告.docx"
    counter = 1
    while output_path.exists():
        output_path = output_dir / f"牙齿病变批量辅助识别报告_{counter:02d}.docx"
        counter += 1

    summary = batch_summary or build_batch_summary(batch_state)
    document = Document()
    document.add_heading("牙齿病变批量辅助识别报告", level=0)
    notice = document.add_paragraph()
    run = notice.add_run("本结果仅供辅助参考，不能替代专业牙科医生诊断。")
    run.bold = True
    document.add_paragraph(f"生成时间：{_summary_value(summary, '生成时间', default='') or ''}")
    document.add_paragraph("报告类型：批量影像检测合并报告")

    document.add_heading("一、批量检测总览", level=1)
    for label in ["图片总数", "成功处理", "处理失败", "检测框总数", "平均置信度", "最高置信度", "涉及类别"]:
        document.add_paragraph(f"{label}：{summary.get(label, '无' if label == '涉及类别' else 0)}")

    document.add_heading("二、类别统计", level=1)
    class_rows = summary.get("类别统计") or []
    if class_rows:
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["类别", "中文名称", "检测框数量", "涉及图片数", "最高置信度", "平均置信度"]
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header
        for row in class_rows:
            cells = table.add_row().cells
            for cell, header in zip(cells, headers):
                cell.text = str(row.get(header, ""))
    else:
        document.add_paragraph("暂无类别统计。")

    document.add_heading("三、重点关注图片", level=1)
    focus_rows = summary.get("重点关注图片") or []
    if focus_rows:
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["排名", "图片名称", "最高类别", "最高置信度", "检测框数量", "关注等级"]
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header
        for row in focus_rows:
            cells = table.add_row().cells
            for cell, header in zip(cells, headers):
                cell.text = str(row.get(header, ""))
    else:
        document.add_paragraph("暂无重点关注图片。")

    document.add_heading("四、逐图检测结果", level=1)
    with TemporaryDirectory() as temp_dir:
        temp_root = Path(temp_dir)
        for index, item in enumerate(batch_state, start=1):
            result = _primary_result(item)
            item_results = _item_results(item)
            name = item.get("name") or item.get("image_name") or f"image_{index:03d}.png"
            display_name = str(item.get("display_name") or name)
            result_detections = [
                det
                for model_result in item_results
                for det in (model_result.get("detections", []) if isinstance(model_result, dict) else [])
                if has_detection_payload(det)
            ]
            detections = [
                det
                for det in (result.get("detections", []) if isinstance(result, dict) else [])
                if has_detection_payload(det)
            ]
            enriched = [enrich_detection_row(det) for det in result_detections]
            confidences = []
            for row in enriched:
                confidence = parse_confidence(row.get("confidence"))
                if confidence is not None:
                    confidences.append(confidence)
            classes = sorted({row.get("中文名称", "") for row in enriched if row.get("中文名称")})

            document.add_heading(f"{index}. {display_name}", level=2)
            if display_name != str(name):
                document.add_paragraph(f"原始文件名：{name}")
            document.add_paragraph(f"使用模型：{_model_name(item)}")
            document.add_paragraph(f"检测框数量：{len(enriched)}")
            document.add_paragraph(f"涉及类别：{'、'.join(classes) if classes else '无'}")
            document.add_paragraph(f"最高置信度：{max(confidences):.4f}" if confidences else "最高置信度：无")
            if len(item_results) > 1:
                document.add_paragraph(f"模型结果组数：{len(item_results)}")
            quality_text = item.get("quality_text")
            if quality_text:
                document.add_paragraph("图像质量提示：")
                _add_paragraphs(document, quality_text)
            annotated = _first_present(result.get("full_annotated"), result.get("annotated")) if isinstance(result, dict) else None
            if annotated is not None:
                document.add_paragraph("检测结果图")
                image_path = _save_temp_image(annotated, temp_root, f"{index:03d}_{_safe_name(name)}.png")
                document.add_picture(str(image_path), width=Inches(5.8))
            image_size = _image_size(result.get("model_input")) or _image_size(result.get("original")) or _image_size(annotated)
            if len(item_results) > 1:
                for model_index, model_result in enumerate(item_results, start=1):
                    model_name = model_result.get("model", f"模型 {model_index}") if isinstance(model_result, dict) else f"模型 {model_index}"
                    model_detections = model_result.get("detections", []) if isinstance(model_result, dict) else []
                    document.add_paragraph(f"模型 {model_index}：{model_name}")
                    model_image = _model_result_image(model_result) if isinstance(model_result, dict) else None
                    if model_image is not None:
                        image_path = _save_temp_image(
                            model_image,
                            temp_root,
                            f"{index:03d}_model_{model_index:02d}_{_safe_name(name)}.png",
                        )
                        document.add_picture(str(image_path), width=Inches(5.8))
                    _add_detection_table(document, model_detections, image_size)
            else:
                _add_detection_table(document, detections, image_size)

            advice = item.get("advice")
            if advice:
                document.add_paragraph("辅助建议：")
                _add_paragraphs(document, advice)

    document.add_heading("五、失败图片列表", level=1)
    failed = summary.get("失败图片") or []
    if failed:
        for item in failed:
            document.add_paragraph(f"- {item}")
    else:
        document.add_paragraph("无失败图片。")

    document.add_heading("六、图例与类别说明", level=1)
    _add_paragraphs(document, legend_markdown())

    document.add_heading("七、安全声明", level=1)
    document.add_paragraph("本系统不适用于最终诊断，不提供治疗方案，不提供药物建议。未检测到目标不代表不存在病变。")
    document.add_paragraph("本结果仅供辅助参考，不能替代专业牙科医生诊断。")

    document.save(output_path)
    return output_path
