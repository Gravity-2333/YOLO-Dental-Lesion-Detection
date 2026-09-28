"""测试脚本 - 检查潜在的bug和边界条件"""
import sys
import json
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

# 测试1: 检查导入
print("测试1: 检查模块导入...")
try:
    from src.dental_detection.assistant import (
        load_settings, normalize_base_url, default_advice,
        _normalize_class_name, CLASS_ADVICE
    )
    from src.dental_detection.inference import Detection
    from src.dental_detection.config import MODEL_REGISTRY, DEFAULT_MODEL_PATH
    from src.dental_detection.assistant import case_dir, ensure_app_dirs
    from src.dental_detection.case_store import list_case_records, load_case_record
    from src.dental_detection.history_store import append_history_records, history_rows, load_history_record
    from src.dental_detection.reporting import SingleReportData, export_batch_docx_report, export_single_docx_report
    from src.dental_detection.batch_summary import build_batch_summary
    from src.dental_detection.record_formatters import format_case_record
    from src.dental_detection.ui_assets import load_workbench_css
    from src.dental_detection.ui_contracts import COMMON_OUTPUT_KEYS
    print("✓ 模块导入成功")
except Exception as e:
    print(f"✗ 模块导入失败: {e}")
    sys.exit(1)


def common_output(outputs, key):
    return outputs[COMMON_OUTPUT_KEYS.index(key)]

# 测试2: 检查配置加载
print("\n测试2: 检查配置加载...")
try:
    settings = load_settings()
    print(f"✓ 配置加载成功")
    print(f"  - Storage dir: {settings.storage_dir}")
    print(f"  - AI enabled: {settings.enabled}")
except Exception as e:
    print(f"✗ 配置加载失败: {e}")

# 测试3: 检查URL规范化
print("\n测试3: 检查URL规范化...")
test_urls = [
    ("api.deepseek.com", "https://api.deepseek.com/v1"),
    ("https://api.openai.com/v1", "https://api.openai.com/v1"),
    ("http://localhost:8000", "http://localhost:8000/v1"),
    ("127.0.0.1:5000", "http://127.0.0.1:5000/v1"),
    ("https://api.example.com/v1/chat/completions", "https://api.example.com/v1"),
    ("https://api.example.com/openai/v1/chat/completions", "https://api.example.com/openai/v1"),
]
for url, expected in test_urls:
    try:
        normalized = normalize_base_url(url)
        assert normalized == expected, f"期望 {expected}，实际 {normalized}"
        print(f"✓ {url} -> {normalized}")
    except Exception as e:
        print(f"✗ {url} 规范化失败: {e}")
        sys.exit(1)

# 测试4: 检查默认建议生成
print("\n测试4: 检查默认建议生成...")
test_detections = [
    [],  # 空检测
    [{"class": "Caries", "confidence": 0.85}],  # 单个检测
    [{"class": "Caries", "confidence": 0.85}, {"class": "Periapical_Lesion", "confidence": 0.72}],  # 多个检测
    [{"class": "Unknown", "confidence": 0.50}],  # 未知类别
]
for i, dets in enumerate(test_detections):
    try:
        advice = default_advice(dets)
        print(f"✓ 测试 {i+1}: 生成建议成功 ({len(advice)} 字符)")
        if i == 0:
            assert "未检测到" in advice, "空检测应包含'未检测到'"
    except Exception as e:
        print(f"✗ 测试 {i+1} 失败: {e}")

# 测试5: 检查类别名称规范化
print("\n测试5: 检查类别名称规范化...")
test_names = [
    ("Periapical_Lesion", "Periapical Lesion"),
    ("Caries", "Caries"),
    ("  Impacted  ", "Impacted"),
]
for original, expected in test_names:
    normalized = _normalize_class_name(original)
    if normalized == expected:
        print(f"✓ '{original}' -> '{normalized}'")
    else:
        print(f"✗ '{original}' -> '{normalized}' (期望: '{expected}')")

# 测试6: 检查模型注册表
print("\n测试6: 检查模型注册表...")
for model_name, model_info in MODEL_REGISTRY.items():
    model_path = Path(model_info["path"])
    if model_path.exists():
        print(f"✓ {model_name}: {model_path.name} 存在")
    else:
        print(f"✗ {model_name}: {model_path} 不存在")

# 测试7: 检查默认模型路径
print("\n测试7: 检查默认模型路径...")
if DEFAULT_MODEL_PATH.exists():
    print(f"✓ 默认模型存在: {DEFAULT_MODEL_PATH}")
else:
    print(f"✗ 默认模型不存在: {DEFAULT_MODEL_PATH}")

# 测试8: 检查Detection dataclass
print("\n测试8: 检查Detection数据类...")
try:
    det = Detection(
        cls_id=0,
        label="Caries",
        confidence=0.85,
        x1=100.0,
        y1=200.0,
        x2=300.0,
        y2=400.0
    )
    row = det.as_row()
    assert "class" in row
    assert "confidence" in row
    assert row["confidence"] == 0.85
    print(f"✓ Detection数据类正常工作")
    print(f"  - Row: {row}")
except Exception as e:
    print(f"✗ Detection数据类测试失败: {e}")

print("\n测试9: 检查空检测Word报告明细...")
try:
    from docx import Document
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    with TemporaryDirectory() as temp_dir:
        report_path = export_single_docx_report(
            SingleReportData(
                image_name="empty.png",
                created_at="2026-06-22T00:00:00",
                model_name="test-model",
                original_image=image,
                model_input_image=image,
                annotated_image=image,
                detections=[],
                advice="测试建议",
                quality_text="测试质量",
                summary={},
                safety_notice="测试声明",
            ),
            Path(temp_dir),
        )
        text = "\n".join(cell.text for table in Document(report_path).tables for row in table.rows for cell in row.cells)
        assert "未检测到目标框" in text, "空检测报告明细表应写明未检测到目标框"
    print("✓ 空检测Word报告明细正常")
except Exception as e:
    print(f"✗ 空检测Word报告测试失败: {e}")
    sys.exit(1)

print("\n测试10: 检查对比模型结果可追溯...")
try:
    from docx import Document
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    detections_a = [{"class": "Caries", "confidence": 0.81, "x1": 1, "y1": 2, "x2": 20, "y2": 30}]
    detections_b = [{"class": "Impacted", "confidence": 0.76, "x1": 10, "y1": 12, "x2": 40, "y2": 50}]
    model_a = {
        "model": "model-a",
        "model_path": "a.pt",
        "detections": detections_a,
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
    }
    model_b = {
        "model": "model-b",
        "model_path": "b.pt",
        "detections": detections_b,
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
    }
    batch_state = [
        {
            "name": "compare.png",
            "display_name": "001 - compare.png",
            "result": model_a,
            "all_results": [model_a, model_b],
            "advice": "测试建议",
            "quality_text": "测试质量",
            "quality_level": "良好",
            "summary": {"模型模式": "对比模型"},
        }
    ]
    with TemporaryDirectory() as temp_dir:
        append_history_records(batch_state, temp_dir, 100)
        rows = history_rows(temp_dir)
        record = load_history_record(rows[0]["记录ID"], temp_dir)
        assert record["detection_count"] == 2, "历史记录应统计两个模型的检测框"
        assert "model-a" in record["model"] and "model-b" in record["model"], "历史记录应保留两个模型名称"
        assert len(record.get("model_results", [])) == 2, "历史记录应保留模型结果明细"

        ensure_app_dirs(temp_dir)
        case_path = case_dir(temp_dir) / "case_20260622_compare.json"
        case_path.write_text(
            json.dumps(
                {
                    "created_at": "2026-06-22T00:00:00",
                    "case_id": "compare-case",
                    "image_name": "compare.png",
                    "summary": {},
                    "model_results": record["model_results"],
                    "detections": detections_a,
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        case_rows = list_case_records(temp_dir)
        assert case_rows[0]["检测数量"] == 2, "病例列表应统计两个模型的检测框"
        loaded_case = load_case_record(temp_dir, case_rows[0]["文件名"])
        assert len(loaded_case["model_results"]) == 2, "病例记录应保留模型结果明细"
        case_detail = format_case_record(loaded_case)
        assert "model-b" in case_detail and "Impacted" in case_detail, "病例详情文本应显示副模型检测框"

        report_path = export_batch_docx_report(batch_state, None, Path(temp_dir) / "reports")
        doc_text = "\n".join(
            paragraph.text for paragraph in Document(report_path).paragraphs
        )
        assert "model-a" in doc_text and "model-b" in doc_text, "批量Word报告应写入两个模型名称"
        assert len(Document(report_path).inline_shapes) >= 3, "批量Word报告应包含主结果图和分模型结果图"

        single_report_path = export_single_docx_report(
            SingleReportData(
                image_name="compare.png",
                created_at="2026-06-22T00:00:00",
                model_name="model-a",
                original_image=image,
                model_input_image=image,
                annotated_image=image,
                detections=detections_a,
                advice="测试建议",
                quality_text="测试质量",
                summary={},
                safety_notice="测试声明",
                model_results=[
                    {"model": "model-a", "model_path": "a.pt", "detections": detections_a, "annotated": image},
                    {"model": "model-b", "model_path": "b.pt", "detections": detections_b, "annotated": image},
                ],
            ),
            Path(temp_dir) / "single_reports",
        )
        single_doc_text = "\n".join(
            cell.text for table in Document(single_report_path).tables for row in table.rows for cell in row.cells
        )
        assert "model-a" in single_doc_text and "model-b" in single_doc_text, "单图Word报告应保留两个模型明细"
        assert len(Document(single_report_path).inline_shapes) >= 5, "单图Word报告应包含每个模型的结果图"
    print("✓ 对比模型结果追溯正常")
except Exception as e:
    print(f"✗ 对比模型追溯测试失败: {e}")
    sys.exit(1)

print("\n测试11: 检查批量同名图片总览不合并...")
try:
    item_a = {
        "name": "same.png",
        "display_name": "001 - same.png",
        "result": {"detections": [{"class": "Caries", "confidence": 0.8}]},
    }
    item_b = {
        "name": "same.png",
        "display_name": "002 - same.png",
        "result": {"detections": [{"class": "Caries", "confidence": 0.7}]},
    }
    summary = build_batch_summary([item_a, item_b])
    caries = next(row for row in summary["类别统计"] if row["类别"] == "Caries")
    assert caries["涉及图片数"] == 2, "同名批量图片应按显示名区分，不能合并为1张"
    focus_names = [row["图片名称"] for row in summary["重点关注图片"]]
    assert "001 - same.png" in focus_names and "002 - same.png" in focus_names, "重点关注图片应显示唯一名称"
    print("✓ 批量同名图片总览正常")
except Exception as e:
    print(f"✗ 批量同名图片总览测试失败: {e}")
    sys.exit(1)

print("\n测试12: 检查对比模式自动修正重复模型...")
try:
    from app import MODEL_MODE_COMPARE, _configured_models, _distinct_compare_model_path, _with_current_defaults
    from src.dental_detection.assistant import AiSettings, DEFAULT_AI_BASE_URL

    resolved = _distinct_compare_model_path(str(DEFAULT_MODEL_PATH), str(DEFAULT_MODEL_PATH))
    assert Path(resolved) != Path(DEFAULT_MODEL_PATH), "重复模型应自动配对另一个项目模型"
    configured = _configured_models(MODEL_MODE_COMPARE, str(DEFAULT_MODEL_PATH), str(DEFAULT_MODEL_PATH))
    assert len(configured) == 2, "自动配对后应返回两个模型"
    assert configured[0][1] != configured[1][1], "自动配对后的模型路径必须不同"

    openai_settings = AiSettings(
        base_url="https://api.openai.com/v1",
        model="gpt-4o-mini",
        key_mode="环境变量",
        api_key="OPENAI_API_KEY",
    )
    preserved = _with_current_defaults(openai_settings, config_exists=True)
    assert preserved.base_url == "https://api.openai.com/v1", "已保存的 OpenAI 配置不应被迁移为 DeepSeek"

    old_default = AiSettings(
        base_url="https://api.openai.com/v1",
        model="gpt-4o-mini",
        key_mode="环境变量",
        api_key="OPENAI_API_KEY",
    )
    migrated = _with_current_defaults(old_default, config_exists=False)
    assert migrated.base_url == DEFAULT_AI_BASE_URL, "首次运行旧默认配置应迁移为当前默认服务"
    print("✓ 对比模式重复模型自动修正正常")
except Exception as e:
    print(f"✗ 对比模式重复模型自动修正测试失败: {e}")
    sys.exit(1)

print("\n测试13: 检查切换图片保留已导出报告路径...")
try:
    from PIL import Image
    from app import select_batch_item

    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "keep-path.png",
        "display_name": "001 - keep-path.png",
        "result": {
            "original": image,
            "model_input": image,
            "annotated": image,
            "table": [],
            "detections": [],
        },
        "advice": "测试建议",
        "quality_text": "测试质量",
        "summary": {},
        "word_report_path": r"C:\tmp\report.docx",
        "zip_report_path": r"C:\tmp\report.zip",
    }
    outputs = select_batch_item("001 - keep-path.png", [item], False)
    word_report_path_index = 21
    zip_report_path_index = 24
    assert outputs[word_report_path_index] == r"C:\tmp\report.docx", "切换图片后应保留 Word 报告路径"
    assert outputs[zip_report_path_index] == r"C:\tmp\report.zip", "切换图片后应保留 ZIP 报告路径"
    print("✓ 切换图片保留报告路径正常")
except Exception as e:
    print(f"✗ 切换图片报告路径测试失败: {e}")
    sys.exit(1)

print("\n测试14: 检查对比模型ZIP导出包含分模型结果图...")
try:
    from PIL import Image
    from app import export_batch_results, export_batch_word_report, export_single_report

    image = Image.new("RGB", (80, 60), "white")
    model_a = {
        "model": "model-a",
        "model_path": "a.pt",
        "detections": [{"class": "Caries", "confidence": 0.81, "x1": 1, "y1": 2, "x2": 20, "y2": 30}],
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
    }
    model_b = {
        "model": "model-b",
        "model_path": "b.pt",
        "detections": [{"class": "Impacted", "confidence": 0.76, "x1": 10, "y1": 12, "x2": 40, "y2": 50}],
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
    }
    batch_state = [
        {
            "name": "compare.png",
            "display_name": "001 - compare.png",
            "result": model_a,
            "all_results": [model_a, model_b],
            "advice": "测试建议",
            "quality_text": "测试质量",
            "quality_level": "良好",
            "summary": {"模型模式": "对比模型"},
        }
    ]
    with TemporaryDirectory() as temp_dir:
        _, _, single_state = export_single_report(batch_state, "001 - compare.png", temp_dir)
        single_zip = Path(single_state[0]["zip_report_path"])
        with zipfile.ZipFile(single_zip) as archive:
            names = archive.namelist()
        assert any("model_01" in name and name.endswith("_result.png") for name in names), "单图ZIP应包含主模型结果图"
        assert any("model_02" in name and name.endswith("_result.png") for name in names), "单图ZIP应包含副模型结果图"

        _, batch_zip_path_text, batch_state = export_batch_results(batch_state, temp_dir)
        batch_zip = Path(batch_state[0]["zip_report_path"])
        assert Path(batch_zip_path_text) == batch_zip, "批量 ZIP 路径框应返回可直接使用的纯路径"
        with zipfile.ZipFile(batch_zip) as archive:
            names = archive.namelist()
        assert any("model_01" in name and name.endswith("_result.png") for name in names), "批量ZIP应包含主模型结果图"
        assert any("model_02" in name and name.endswith("_result.png") for name in names), "批量ZIP应包含副模型结果图"

        _, batch_word_path_text, batch_state = export_batch_word_report(batch_state, temp_dir)
        assert Path(batch_word_path_text).is_file(), "批量 Word 路径框应返回可直接使用的纯路径"
    print("✓ 对比模型ZIP分模型结果图正常")
except Exception as e:
    print(f"✗ 对比模型ZIP分模型结果图测试失败: {e}")
    sys.exit(1)

print("\n测试15: 检查历史记录保留ZIP报告路径...")
try:
    from src.dental_detection.record_formatters import format_history_record

    with TemporaryDirectory() as temp_dir:
        item = {
            "name": "zip-only.png",
            "display_name": "001 - zip-only.png",
            "result": {"model": "model-a", "detections": []},
            "zip_report_path": r"C:\tmp\zip-only-report.zip",
        }
        append_history_records([item], temp_dir, 100)
        rows = history_rows(temp_dir)
        record = load_history_record(rows[0]["记录ID"], temp_dir)
        assert record["report_path"] == r"C:\tmp\zip-only-report.zip", "历史记录应在没有Word路径时保留ZIP报告路径"
        detail = format_history_record(record)
        assert "报告文件：zip-only-report.zip" in detail, "历史详情应显示ZIP报告文件名"
        assert r"C:\tmp" not in detail, "历史详情不应显示ZIP报告绝对目录"

        latest_item = {
            "name": "latest-report.png",
            "display_name": "001 - latest-report.png",
            "result": {"model": "model-a", "detections": []},
            "word_report_path": r"C:\tmp\old-word-report.docx",
            "zip_report_path": r"C:\tmp\latest-report.zip",
            "report_path": r"C:\tmp\latest-report.zip",
        }
        append_history_records([latest_item], temp_dir, 100)
        latest_rows = history_rows(temp_dir)
        latest_record = load_history_record(latest_rows[0]["记录ID"], temp_dir)
        assert latest_record["report_path"] == r"C:\tmp\latest-report.zip", "历史记录应优先保留最近一次导出的报告路径"
    print("✓ 历史记录ZIP报告路径正常")
except Exception as e:
    print(f"✗ 历史记录ZIP报告路径测试失败: {e}")
    sys.exit(1)

print("\n测试16: 检查历史/病例兼容百分号置信度和中文字段...")
try:
    from src.dental_detection.result_levels import enrich_detection_row, parse_confidence

    assert parse_confidence("81%") == 0.81, "百分号置信度应转换为0-1小数"
    assert parse_confidence(81) == 0.81, "0-100置信度数值应转换为0-1小数"
    enriched = enrich_detection_row({"类别": "Periapical_Lesion", "置信度": "76%", "x1": 1, "y1": 2, "x2": 20, "y2": 30})
    assert enriched["class"] == "Periapical_Lesion", "中文类别字段应被识别"
    assert enriched["confidence"] == 0.76, "中文置信度字段应被识别"
    assert enriched["关注等级"] == "重点关注", "百分号置信度应参与关注等级判断"

    with TemporaryDirectory() as temp_dir:
        legacy_item = {
            "name": "legacy.png",
            "result": {
                "model": "legacy-model",
                "detections": [
                    {"类别": "Caries", "置信度": "81%", "x1": 1, "y1": 2, "x2": 20, "y2": 30},
                ],
            },
        }
        append_history_records([legacy_item], temp_dir, 100)
        rows = history_rows(temp_dir)
        assert rows[0]["最高置信度"] == "0.81", "历史表应正确显示百分号置信度"
        assert rows[0]["关注等级"] == "重点关注", "历史表应正确计算百分号关注等级"

        ensure_app_dirs(temp_dir)
        case_path = case_dir(temp_dir) / "case_20260622_legacy.json"
        case_path.write_text(
            json.dumps(
                {
                    "created_at": "2026-06-22T00:00:00",
                    "case_id": "legacy-case",
                    "image_name": "legacy.png",
                    "detections": [{"类别": "Impacted", "置信度": 76, "x1": 1, "y1": 2, "x2": 20, "y2": 30}],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        case_rows = list_case_records(temp_dir)
        assert case_rows[0]["最高置信度"] == 0.76, "病例列表应正确统计0-100置信度"
        assert case_rows[0]["关注等级"] == "重点关注", "病例列表应正确计算0-100关注等级"
    print("✓ 百分号置信度和中文字段兼容正常")
except Exception as e:
    print(f"✗ 百分号置信度和中文字段兼容测试失败: {e}")
    sys.exit(1)

print("\n测试17: 检查病例记录兼容空模型明细回退顶层检测框...")
try:
    with TemporaryDirectory() as temp_dir:
        ensure_app_dirs(temp_dir)
        case_path = case_dir(temp_dir) / "case_20260705_fallback.json"
        case_path.write_text(
            json.dumps(
                {
                    "created_at": "2026-07-05T00:00:00",
                    "case_id": "fallback-case",
                    "image_name": "fallback.png",
                    "model_results": [{"model": "legacy-model", "detections": []}],
                    "detections": [{"class": "Caries", "confidence": 0.83, "x1": 1, "y1": 2, "x2": 20, "y2": 30}],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        case_rows = list_case_records(temp_dir)
        assert case_rows[0]["检测数量"] == 1, "空模型明细不应掩盖顶层 detections"
        loaded_case = load_case_record(temp_dir, case_rows[0]["文件名"])
        detail = format_case_record(loaded_case)
        assert "Caries" in detail and "龋齿" in detail, "病例详情应回退显示顶层检测框"
    print("✓ 病例空模型明细兼容正常")
except Exception as e:
    print(f"✗ 病例空模型明细兼容测试失败: {e}")
    sys.exit(1)

print("\n测试18: 检查对比模型建议覆盖副模型检测框...")
try:
    from app import _advice_detections
    from src.dental_detection.assistant import default_advice

    merged = _advice_detections(
        [
            {"model": "model-a", "detections": []},
            {
                "model": "model-b",
                "detections": [{"class": "Impacted", "confidence": 0.82, "x1": 1, "y1": 2, "x2": 20, "y2": 30}],
            },
        ]
    )
    assert len(merged) == 1, "建议输入应包含副模型独有检测框"
    assert merged[0]["model"] == "model-b" and merged[0]["模型"] == "model-b", "建议输入应保留检测框来源模型"
    advice = default_advice(merged)
    assert "阻生牙" in advice, "默认建议应覆盖副模型检出的类别"
    print("✓ 对比模型建议覆盖副模型检测框正常")
except Exception as e:
    print(f"✗ 对比模型建议覆盖副模型检测框测试失败: {e}")
    sys.exit(1)

print("\n测试19: 检查批量摘要和默认建议兼容中文检测字段...")
try:
    chinese_detection = {"类别": "Periapical_Lesion", "置信度": "82%", "x1": 1, "y1": 2, "x2": 20, "y2": 30}
    summary = build_batch_summary(
        [
            {
                "name": "legacy-cn.png",
                "display_name": "001 - legacy-cn.png",
                "result": {"detections": [chinese_detection]},
            }
        ]
    )
    assert summary["检测框总数"] == 1, "中文字段检测框应参与批量统计"
    assert summary["最高置信度"] == 0.82, "中文置信度字段应参与最高置信度统计"
    assert summary["涉及类别"] == "根尖周病变", "中文字段类别应正确映射到显示名称"
    assert summary["重点关注图片"][0]["原始类别"] == "Periapical_Lesion", "重点关注图片应保留规范类别"
    advice = default_advice([chinese_detection])
    assert "根尖周" in advice and "0.82" in advice, "默认建议应兼容中文类别和百分号置信度"
    print("✓ 批量摘要和默认建议中文字段兼容正常")
except Exception as e:
    print(f"✗ 批量摘要和默认建议中文字段兼容测试失败: {e}")
    sys.exit(1)

print("\n测试20: 检查不存在的历史报告文件不会显示为可下载文件...")
try:
    from app import _file_component_output, _remember_allowed_file_root

    with TemporaryDirectory() as temp_dir:
        missing_path = Path(temp_dir) / "missing-report.docx"
        _remember_allowed_file_root(temp_dir)
        output = _file_component_output(missing_path)
        assert output["visible"] is False, "不存在的报告文件不应显示下载组件"
        assert output["value"] is None, "不存在的报告文件不应作为 File 组件值"
    print("✓ 不存在报告文件下载入口隐藏正常")
except Exception as e:
    print(f"✗ 不存在报告文件下载入口测试失败: {e}")
    sys.exit(1)

print("\n测试21: 检查受管目录迁移到自身子目录不递归搬动新目录...")
try:
    from src.dental_detection.assistant import migrate_storage

    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        old_cases = root / "cases"
        old_cases.mkdir()
        (old_cases / "case_legacy.json").write_text("{}", encoding="utf-8")
        new_storage = old_cases / "nested_storage"
        migrate_storage(str(old_cases), str(new_storage))
        migrated_case = new_storage / "cases" / "case_legacy.json"
        recursive_target = new_storage / "cases" / "nested_storage"
        assert migrated_case.exists(), "旧 cases 根目录里的病例文件应迁入新存储目录的 cases 子目录"
        assert not recursive_target.exists(), "迁移时不应把新存储目录再搬进自身子目录"
    print("✓ 受管目录迁移到自身子目录正常")
except Exception as e:
    print(f"✗ 受管目录迁移测试失败: {e}")
    sys.exit(1)

print("\n测试22: 检查检测框标签贴边时不会越出图像...")
try:
    from PIL import Image, ImageDraw
    from src.dental_detection.visualization import draw_detections_with_filter, label_box_layout, label_font_for_image

    left, top, right, bottom, text_x, text_y = label_box_layout(
        image_width=80,
        image_height=60,
        anchor_x=76,
        anchor_y=4,
        text_width=42,
        text_height=12,
    )
    assert 0 <= left <= right <= 79, "标签背景应被限制在图像宽度内"
    assert 0 <= top <= bottom <= 59, "标签背景应被限制在图像高度内"
    assert text_x >= left and text_y >= top, "文字起点应位于标签背景内部"

    image = Image.new("RGB", (80, 60), "white")
    annotated = draw_detections_with_filter(
        image,
        [{"class": "Caries", "confidence": 0.91, "x1": 76, "y1": 4, "x2": 79, "y2": 30}],
    )
    assert annotated.size == image.size, "重绘结果不应改变图像尺寸"
    large_font = label_font_for_image((1800, 900))
    text_box = ImageDraw.Draw(Image.new("RGB", (1, 1))).textbbox((0, 0), "Caries 0.85", font=large_font)
    assert text_box[3] - text_box[1] >= 18, "检测结果标签字号应明显大于 PIL 默认小字体"
    print("✓ 检测框标签贴边布局正常")
except Exception as e:
    print(f"✗ 检测框标签贴边布局测试失败: {e}")
    sys.exit(1)

print("\n测试23: 检查模型文件不存在时优先提示路径缺失...")
try:
    from src.dental_detection.error_messages import friendly_error_message

    message = friendly_error_message("model file not found: C:/missing/best.pt", "模型文件不存在")
    assert "相关文件不存在" in message, "模型缺失应优先提示文件不存在，而不是泛化为加载失败"
    assert "模型文件无法加载" not in message, "缺失文件不应被误判为权重损坏"
    print("✓ 模型缺失错误提示正常")
except Exception as e:
    print(f"✗ 模型缺失错误提示测试失败: {e}")
    sys.exit(1)

print("\n测试24: 检查模型扫描包含 mlpackage 目录模型...")
try:
    from src.dental_detection.model_files import is_supported_model_artifact, scan_model_files, supported_suffix_text

    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        package = root / "exported_model.mlpackage"
        package.mkdir()
        fake_package_file = root / "fake_package.mlpackage"
        fake_package_file.write_bytes(b"not a package directory")
        choices = scan_model_files(root)
        assert any(Path(value).name == "exported_model.mlpackage" for _, value in choices), (
            "扫描模型目录时应把 .mlpackage 目录作为可选模型"
        )
        assert not any(Path(value).name == "fake_package.mlpackage" for _, value in choices), (
            ".mlpackage 普通文件不应被误认为可选模型包"
        )
        assert is_supported_model_artifact(package), ".mlpackage 目录应被视为支持的模型包"
        assert not is_supported_model_artifact(fake_package_file), ".mlpackage 普通文件不应通过模型包校验"
        assert ".mlpackage" in supported_suffix_text(), "支持格式提示仍应包含 .mlpackage 模型包"

        direct_choices = scan_model_files(package)
        assert len(direct_choices) == 1 and Path(direct_choices[0][1]).name == "exported_model.mlpackage", (
            "直接选择 .mlpackage 目录时应返回该模型"
        )
    print("✓ mlpackage 目录模型扫描正常")
except Exception as e:
    print(f"✗ mlpackage 目录模型扫描测试失败: {e}")
    sys.exit(1)

print("\n测试25: 检查示例图片加载后不依赖未关闭文件句柄...")
try:
    from app import EXAMPLE_DIR, load_demo_example

    example_path = next(EXAMPLE_DIR.glob("*.png"))
    image, info = load_demo_example(str(example_path))
    assert image.mode == "RGB", "示例图片应加载为 RGB"
    assert image.size[0] > 0 and image.size[1] > 0, "示例图片尺寸应有效"
    assert str(info).strip(), "示例图片应返回说明文本"
    with TemporaryDirectory() as temp_dir:
        output_path = Path(temp_dir) / "example_copy.png"
        image.save(output_path)
        assert output_path.exists(), "示例图片对象应可独立保存，不依赖打开的源文件"
    print("✓ 示例图片加载句柄处理正常")
except Exception as e:
    print(f"✗ 示例图片加载测试失败: {e}")
    sys.exit(1)

print("\n测试26: 检查未选择历史记录时删除提示不会覆盖详情...")
try:
    from app import delete_selected_history_record

    with TemporaryDirectory() as temp_dir:
        _, _, detail, feedback = delete_selected_history_record("", temp_dir)
        assert "请选择一条检测历史" in detail, "未选择时历史详情应保持正常空状态"
        assert "请选择要删除的历史记录" in feedback, "未选择删除应把提示放在反馈区域"
    print("✓ 历史删除空选择提示正常")
except Exception as e:
    print(f"✗ 历史删除空选择提示测试失败: {e}")
    sys.exit(1)

print("\n测试27: 检查空模型明细不会掩盖主结果检测框...")
try:
    from docx import Document
    from PIL import Image
    from app import _item_results, export_batch_results
    from src.dental_detection.record_formatters import format_history_record

    image = Image.new("RGB", (80, 60), "white")
    primary = {
        "model": "primary-model",
        "model_path": "primary.pt",
        "detections": [{"class": "Caries", "confidence": 0.84, "x1": 1, "y1": 2, "x2": 20, "y2": 30}],
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
        "table": [],
    }
    item = {
        "name": "stale-results.png",
        "display_name": "001 - stale-results.png",
        "result": primary,
        "all_results": [{"model": "stale-empty-model", "detections": []}],
        "advice": "测试建议",
        "quality_text": "测试质量",
        "quality_level": "良好",
        "summary": {"模型模式": "单模型"},
    }
    assert _item_results(item)[0]["model"] == "primary-model", "空 all_results 不应掩盖主结果"
    summary = build_batch_summary([item])
    assert summary["检测框总数"] == 1, "批量摘要应回退统计主结果检测框"
    with TemporaryDirectory() as temp_dir:
        append_history_records([item], temp_dir, 100)
        rows = history_rows(temp_dir)
        record = load_history_record(rows[0]["记录ID"], temp_dir)
        assert record["detection_count"] == 1, "历史记录应回退保存主结果检测框"
        assert "龋齿" in format_history_record(record), "历史详情应显示主结果检测框"

        report_path = export_batch_docx_report([item], summary, Path(temp_dir) / "word")
        doc = Document(report_path)
        doc_text = "\n".join(
            [paragraph.text for paragraph in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
        assert "primary-model" in doc_text and "Caries" in doc_text, "批量 Word 应回退显示主结果检测框"

        _, _, exported_state = export_batch_results([item], temp_dir)
        batch_zip = Path(exported_state[0]["zip_report_path"])
        with zipfile.ZipFile(batch_zip) as archive:
            detections_json = json.loads(archive.read("detections.json").decode("utf-8"))
        exported_item = detections_json["items"][0]
        assert exported_item["models"][0]["model"] == "primary-model", "批量 ZIP 应回退主模型名称"
        assert exported_item["models"][0]["detections"][0]["class"] == "Caries", "批量 ZIP 应回退主检测框"
    print("✓ 空模型明细回退主结果正常")
except Exception as e:
    print(f"✗ 空模型明细回退主结果测试失败: {e}")
    sys.exit(1)

print("\n测试28: 检查测试AI接口不会保存配置...")
try:
    import app

    original_test_chat = app.test_chat_completion
    original_save_runtime = app._save_runtime_settings
    try:
        app.test_chat_completion = lambda settings: "测试成功：OK"

        def fail_if_saved(*args, **kwargs):
            raise AssertionError("测试接口不应保存运行时设置")

        app._save_runtime_settings = fail_if_saved
        result = app.test_ai_settings(
            True,
            "https://api.example.com/v1",
            "example-model",
            "直接 Key 值",
            "EXAMPLE_API_KEY",
            "sk-hidden",
            "sk-visible",
            True,
            False,
            True,
            "E:/tmp/should-not-persist",
            "测试 prompt",
            "简洁版",
        )
        assert result == "测试成功：OK", "测试接口应返回测试结果"

        disabled_result = app.test_ai_settings(
            False,
            "https://api.example.com/v1",
            "example-model",
            "直接 Key 值",
            "EXAMPLE_API_KEY",
            "sk-hidden",
            "sk-visible",
            True,
            False,
            True,
            "E:/tmp/should-not-persist",
            "测试 prompt",
            "简洁版",
        )
        assert "AI 功能未开启" in disabled_result, "AI 关闭时应只提示，不保存配置"
    finally:
        app.test_chat_completion = original_test_chat
        app._save_runtime_settings = original_save_runtime
    print("✓ 测试AI接口不保存配置正常")
except Exception as e:
    print(f"✗ 测试AI接口保存隔离测试失败: {e}")
    sys.exit(1)

print("\n测试29: 检查无检测主结果不会被陈旧空明细覆盖模型名...")
try:
    from PIL import Image
    from app import _item_results, export_single_report

    image = Image.new("RGB", (80, 60), "white")
    primary = {
        "model": "primary-empty-model",
        "model_path": "primary-empty.pt",
        "detections": [],
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
        "table": [],
    }
    item = {
        "name": "empty-with-primary.png",
        "display_name": "001 - empty-with-primary.png",
        "result": primary,
        "all_results": [{"model": "stale-empty-model", "detections": []}],
        "advice": "未检测到目标框。",
        "quality_text": "测试质量",
        "quality_level": "良好",
        "summary": {"模型模式": "单模型"},
    }
    assert _item_results(item)[0]["model"] == "primary-empty-model", "无检测主结果也应优先保留真实模型名"
    with TemporaryDirectory() as temp_dir:
        append_history_records([item], temp_dir, 100)
        rows = history_rows(temp_dir)
        record = load_history_record(rows[0]["记录ID"], temp_dir)
        assert record["model"] == "primary-empty-model", "历史记录不应写入陈旧空明细模型名"

        _, _, exported_state = export_single_report([item], "001 - empty-with-primary.png", temp_dir)
        single_zip = Path(exported_state[0]["zip_report_path"])
        with zipfile.ZipFile(single_zip) as archive:
            detections_json = json.loads(archive.read("detections.json").decode("utf-8"))
        assert detections_json["report"]["model"] == "primary-empty-model", "单图 ZIP 报告模型名应来自主结果"
        assert detections_json["models"][0]["model"] == "primary-empty-model", "单图 ZIP 模型明细应来自主结果"
    print("✓ 无检测主结果模型名保留正常")
except Exception as e:
    print(f"✗ 无检测主结果模型名保留测试失败: {e}")
    sys.exit(1)

print("\n测试30: 检查下载结果图文件名过滤控制字符...")
try:
    from PIL import Image
    from src.dental_detection.visualization import save_result_image

    image = Image.new("RGB", (16, 16), "white")
    with TemporaryDirectory() as temp_dir:
        path = save_result_image(image, temp_dir, "牙片\n\tCON:01?.png")
        assert path.exists(), "结果图应能保存成功"
        assert "\n" not in path.name and "\t" not in path.name, "文件名不应包含控制字符"
        assert ":" not in path.name and "?" not in path.name, "文件名不应包含 Windows 禁用字符"
        assert "牙片" in path.name, "安全清理后应尽量保留中文可读信息"
    print("✓ 下载结果图文件名清理正常")
except Exception as e:
    print(f"✗ 下载结果图文件名清理测试失败: {e}")
    sys.exit(1)

print("\n测试31: 检查紧凑工具行按钮底部对齐...")
try:
    css_text = load_workbench_css()
    compact_rule = ".compact-row > .secondary-action,\n.compact-row > .primary-action,\n.compact-row > button"
    assert compact_rule in css_text, "应保留 compact-row 按钮对齐规则"
    rule_start = css_text.index(compact_rule)
    rule_end = css_text.index("}", rule_start)
    rule_body = css_text[rule_start:rule_end]
    assert "align-self: end" in rule_body, "compact-row 内按钮应与输入框底部对齐"
    assert "align-self: center" not in rule_body, "compact-row 内按钮不应垂直居中导致偏上"
    path_rule = (
        ".path-row > .icon-action,\n.path-row > .secondary-action,\n"
        ".model-row > .secondary-action,\n.chat-input-row > .primary-action"
    )
    assert path_rule in css_text, "路径行图标按钮和命令按钮应使用同一底部补偿规则"
    path_rule_start = css_text.index(path_rule)
    path_rule_end = css_text.index("}", path_rule_start)
    assert "margin-bottom: 10px" in css_text[path_rule_start:path_rule_end], (
        "路径选择按钮应与输入框和相邻命令按钮保持同一基线"
    )
    print("✓ 紧凑工具行按钮底部对齐正常")
except Exception as e:
    print(f"✗ 紧凑工具行按钮对齐测试失败: {e}")
    sys.exit(1)

print("\n测试32: 检查病例筛选短输入框不换行截断...")
try:
    project_root = Path(__file__).parent
    cases_page_text = (project_root / "src" / "dental_detection" / "ui_cases_page.py").read_text(encoding="utf-8")
    css_text = load_workbench_css()
    for label in ["病例编号 / 备注名称", "搜索病例", "开始日期", "结束日期"]:
        assert f'label="{label}"' in cases_page_text, f"{label} 输入框应存在"
        label_index = cases_page_text.index(f'label="{label}"')
        snippet = cases_page_text[label_index : label_index + 220]
        assert "lines=1" in snippet and "max_lines=1" in snippet, f"{label} 应声明为单行输入框"
        assert f'textarea[aria-label="{label}"]' in css_text, f"{label} 应有单行样式兜底"
    short_input_rule = 'textarea[aria-label="病例编号 / 备注名称"],\ntextarea[aria-label="搜索病例"],'
    rule_start = css_text.index(short_input_rule)
    rule_end = css_text.index("}", rule_start)
    rule_body = css_text[rule_start:rule_end]
    assert "white-space: nowrap" in rule_body and "text-overflow: ellipsis" in rule_body, "短输入框占位文字应单行省略"
    print("✓ 病例筛选短输入框单行显示正常")
except Exception as e:
    print(f"✗ 病例筛选短输入框测试失败: {e}")
    sys.exit(1)

print("\n测试33: 检查路径输入框长文本省略规则覆盖 input...")
try:
    css_text = load_workbench_css()
    path_rule_anchor = 'textarea[aria-label="报告路径"],'
    rule_start = css_text.index(path_rule_anchor)
    rule_end = css_text.index("}", rule_start)
    rule_body = css_text[rule_start:rule_end]
    for label in [
        "报告路径",
        "导出路径",
        "批量导出路径",
        "批量 Word 报告路径",
        "Word 报告路径",
        "结果图路径",
        "病例报告路径",
        "模型目录",
        "存储目录",
        "主模型路径",
        "对比模型路径",
    ]:
        assert f'textarea[aria-label="{label}"]' in rule_body, f"{label} textarea 应有长文本省略规则"
        assert f'input[aria-label="{label}"]' in rule_body, f"{label} input 应有长文本省略规则"
    assert ".path-row input" in rule_body, "路径行运行时 input 应有长文本省略规则"
    assert ".settings-card input:not" in rule_body, "设置卡片单行 input 应有长文本省略规则"
    assert "text-overflow: ellipsis" in rule_body and "white-space: nowrap" in rule_body, "路径输入框应单行省略"
    print("✓ 路径输入框长文本省略规则正常")
except Exception as e:
    print(f"✗ 路径输入框省略规则测试失败: {e}")
    sys.exit(1)

print("\n测试34: 检查前端更多菜单脚本等待 body 可用...")
try:
    js_text = (Path(__file__).parent / "assets" / "workbench.js").read_text(encoding="utf-8")
    assert "const start = () =>" in js_text, "前端脚本应封装启动逻辑"
    assert "if (!document.body)" in js_text and "requestAnimationFrame(start)" in js_text, (
        "MutationObserver 注册前应等待 document.body 可用"
    )
    assert "scheduleMenuLabeling" in js_text and "labelingScheduled" in js_text, (
        "大量组件切换时应按动画帧合并菜单标记任务，避免重复扫描整个页面"
    )
    assert "new MutationObserver(scheduleMenuLabeling)" in js_text, (
        "DOM 观察器不应在每批变更中同步扫描全部菜单"
    )
    assert "DOMContentLoaded" in js_text and "{ once: true }" in js_text, "脚本应兼容提前注入场景"
    print("✓ 前端更多菜单脚本加载时机正常")
except Exception as e:
    print(f"✗ 前端更多菜单脚本测试失败: {e}")
    sys.exit(1)

print("\n测试35: 检查批量部分失败不会污染首张图片建议...")
try:
    import app
    from PIL import Image

    class FakeFile:
        def __init__(self, name: str):
            self.name = name

    original_detect = app._detect_model_path
    try:
        image = Image.new("RGB", (80, 60), "white")

        def fake_detect(model_name, model_path, image_path, use_clahe, conf, iou, device):
            if "bad" in str(image_path):
                raise RuntimeError("broken image")
            detections = [
                {
                    "class": "Caries",
                    "confidence": 0.82,
                    "x1": 10,
                    "y1": 10,
                    "x2": 40,
                    "y2": 40,
                }
            ]
            return {
                "model": model_name,
                "model_path": str(model_path),
                "original": image,
                "model_input": image,
                "annotated": image,
                "full_annotated": image,
                "detections": detections,
                "table": app._table_from_records(detections),
                "class_names": {"0": "Caries"},
            }

        app._detect_model_path = fake_detect
        with TemporaryDirectory() as temp_dir:
            model_path = Path(temp_dir) / "model.pt"
            model_path.write_bytes(b"fake")
            outputs = app.run_batch_detection(
                [FakeFile("ok.png"), FakeFile("bad.png")],
                app.MODEL_MODE_SINGLE,
                "personal-self",
                str(model_path),
                str(model_path),
                0.25,
                0.7,
                "cpu",
                False,
                False,
                False,
                False,
                "https://api.example.com/v1",
                "example-model",
                "环境变量",
                "EXAMPLE_API_KEY",
                "",
                "",
                False,
                False,
                False,
                temp_dir,
                "",
                "简洁版",
                False,
                100,
            )
        advice_text = common_output(outputs, "advice")
        overview_update = common_output(outputs, "batch_overview")
        overview_html = overview_update.get("value", "") if isinstance(overview_update, dict) else overview_update
        batch_state = common_output(outputs, "batch_state")
        assert "bad.png" not in advice_text and "处理失败" not in advice_text, (
            "首张成功图片的建议不应混入批量失败清单"
        )
        assert "bad.png" in overview_html and "失败图片" in overview_html, "失败图片应在批量总览中独立展示"
        assert batch_state[0].get("batch_errors"), "批量失败信息应保留在状态里供导出使用"
    finally:
        app._detect_model_path = original_detect
    print("✓ 批量失败信息独立展示正常")
except Exception as e:
    print(f"✗ 批量失败信息污染建议测试失败: {e}")
    sys.exit(1)

print("\n测试36: 检查带图片产物的空模型明细不覆盖主检测框...")
try:
    import app
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    primary = {
        "model": "primary-with-box",
        "detections": [{"class": "Caries", "confidence": 0.81, "x1": 1, "y1": 2, "x2": 20, "y2": 30}],
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
    }
    stale_empty_with_artifacts = {
        "model": "stale-empty-with-image",
        "detections": [],
        "original": image,
        "model_input": image,
        "annotated": image,
        "full_annotated": image,
    }
    item = {
        "name": "artifact-stale.png",
        "display_name": "001 - artifact-stale.png",
        "result": primary,
        "all_results": [stale_empty_with_artifacts],
    }
    resolved = app._item_results(item)
    assert resolved[0]["model"] == "primary-with-box", "有图但空检测的陈旧明细不应覆盖主检测框"
    assert build_batch_summary([item])["检测框总数"] == 1, "批量摘要应统计主检测框"

    no_detection_item = {
        "name": "compare-empty.png",
        "result": {
            "model": "primary-empty",
            "detections": [],
            "original": image,
            "model_input": image,
            "annotated": image,
        },
        "all_results": [
            {"model": "empty-a", "detections": [], "annotated": image},
            {"model": "empty-b", "detections": [], "annotated": image},
        ],
    }
    no_detection_results = app._item_results(no_detection_item)
    assert [row["model"] for row in no_detection_results] == ["empty-a", "empty-b"], (
        "真正无检测的对比结果仍应保留分模型明细"
    )
    print("✓ 带图片产物的空模型明细回退正常")
except Exception as e:
    print(f"✗ 带图片产物空模型明细测试失败: {e}")
    sys.exit(1)

print("\n测试37: 检查伪装成模型文件的目录会被提前拦截...")
try:
    import app

    with TemporaryDirectory() as temp_dir:
        fake_pt_dir = Path(temp_dir) / "fake_model.pt"
        fake_pt_dir.mkdir()
        try:
            app._validate_model_files([("主模型", fake_pt_dir)])
            raise AssertionError(".pt 目录不应通过模型文件校验")
        except Exception as exc:
            assert "模型路径无效" in str(exc) or "not a file" in str(exc), "应提示模型路径无效"

        feedback = app.test_model_file(str(fake_pt_dir), str(fake_pt_dir), app.MODEL_MODE_SINGLE)
        assert "路径不是可加载的模型文件" in feedback, "测试模型应明确提示目录不是模型文件"

        package_dir = Path(temp_dir) / "exported_model.mlpackage"
        package_dir.mkdir()
        app._validate_model_files([("CoreML 模型包", package_dir)])
        fake_package_file = Path(temp_dir) / "fake_package.mlpackage"
        fake_package_file.write_bytes(b"not a package directory")
        try:
            app._validate_model_files([("伪模型包", fake_package_file)])
            raise AssertionError(".mlpackage 普通文件不应通过模型包校验")
        except Exception as exc:
            assert "模型路径无效" in str(exc) or "not a file or supported model package" in str(exc), (
                ".mlpackage 普通文件应提示不是可加载模型文件或模型包"
            )
        original_yolo = app.YOLO
        try:
            class FakeYOLO:
                names = {0: "Caries"}

                def __init__(self, path):
                    self.path = path

            app.YOLO = FakeYOLO
            package_feedback = app.test_model_file(str(package_dir), str(package_dir), app.MODEL_MODE_SINGLE)
            assert "文件后缀不在支持列表中" not in package_feedback, (
                "测试模型按钮不应把已支持的 .mlpackage 目录误报为后缀不支持"
            )
            assert "可加载" in package_feedback, "测试模型按钮应继续走到 .mlpackage 加载检查"
        finally:
            app.YOLO = original_yolo
    print("✓ 模型路径文件/目录类型校验正常")
except Exception as e:
    print(f"✗ 模型路径文件/目录类型测试失败: {e}")
    sys.exit(1)

print("\n测试38: 检查坏编码配置、病例和历史文件不会拖垮主流程...")
try:
    import src.dental_detection.assistant as assistant_module
    from src.dental_detection.history_store import history_file

    original_config_path = assistant_module.CONFIG_PATH
    try:
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            bad_settings = root / "settings.json"
            bad_settings.write_bytes(b"\xff\xfe\xff")
            assistant_module.CONFIG_PATH = bad_settings
            settings = assistant_module.load_settings()
            assert settings.model, "坏编码 settings.json 应回退默认设置"
            assert not bad_settings.exists(), "坏编码 settings.json 应被移走备份"
            assert list(root.glob("settings.*.corrupt.json")), "坏编码 settings.json 应保留 corrupt 备份"

        with TemporaryDirectory() as temp_dir:
            ensure_app_dirs(temp_dir)
            bad_case = case_dir(temp_dir) / "case_bad_encoding.json"
            bad_case.write_bytes(b"\xff\xfe\xff")
            case_rows = list_case_records(temp_dir)
            assert case_rows and case_rows[0]["病例编号"] == "损坏病例文件", "坏编码病例文件应显示为损坏记录"
            detail = app.load_case_record("任意时间 | 病例 | 图片 | case_bad_encoding.json", temp_dir)
            assert "错误" in detail or "病例文件" in detail, "坏编码病例详情应返回错误文本而不是抛出异常"

        with TemporaryDirectory() as temp_dir:
            ensure_app_dirs(temp_dir)
            history_file(temp_dir).write_bytes(b"\xff\xfe\xff")
            assert history_rows(temp_dir) == [], "坏编码历史文件应被当作空历史处理"
    finally:
        assistant_module.CONFIG_PATH = original_config_path
    print("✓ 坏编码文件容错正常")
except Exception as e:
    print(f"✗ 坏编码文件容错测试失败: {e}")
    sys.exit(1)

print("\n测试39: 检查空白存储目录不会误建相对目录...")
try:
    from src.dental_detection.assistant import APP_HOME, storage_root

    assert storage_root("   ") == APP_HOME, "纯空白存储目录应回退默认数据目录"
    with TemporaryDirectory() as temp_dir:
        normalized = storage_root(f"  {temp_dir}  ")
        assert normalized == Path(temp_dir), "存储目录前后空格应被清理"
    print("✓ 存储目录空白输入归一正常")
except Exception as e:
    print(f"✗ 存储目录空白输入测试失败: {e}")
    sys.exit(1)

print("\n测试40: 检查历史详情兼容仅含模型明细的旧记录...")
try:
    from src.dental_detection.record_formatters import format_history_record

    legacy_history = {
        "created_at": "2026-07-06T10:00:00",
        "image_name": "legacy-history.png",
        "display_name": "001 - legacy-history.png",
        "model": "legacy-model",
        "detection_count": 1,
        "classes": ["龋齿"],
        "max_confidence": 0.83,
        "level": "重点关注",
        "detections": [],
        "model_results": [
            {
                "model": "detail-model",
                "detection_count": 1,
                "detections": [
                    {
                        "class": "Caries",
                        "confidence": 0.83,
                        "x1": 1,
                        "y1": 2,
                        "x2": 30,
                        "y2": 40,
                    }
                ],
            }
        ],
    }
    detail = format_history_record(legacy_history)
    assert "龋齿" in detail and "detail-model" in detail, "历史详情应从模型明细回退显示检测框"
    assert "无检测框" not in detail, "模型明细存在检测框时不应显示无检测框"
    print("✓ 历史详情模型明细回退正常")
except Exception as e:
    print(f"✗ 历史详情模型明细回退测试失败: {e}")
    sys.exit(1)

print("\n测试41: 检查单图导出路径不会误同步到同名兄弟项...")
try:
    from PIL import Image
    import app

    image = Image.new("RGB", (80, 60), "white")

    def make_item(model_name: str):
        return {
            "name": "same.png",
            "result": {
                "model": model_name,
                "detections": [],
                "original": image,
                "model_input": image,
                "annotated": image,
                "full_annotated": image,
                "table": [],
            },
            "advice": "测试建议",
            "quality_text": "测试质量",
            "quality_level": "良好",
            "summary": {},
        }

    batch_state = [make_item("model-a"), make_item("model-b")]
    with TemporaryDirectory() as temp_dir:
        _, _, exported_state = app.export_single_report(batch_state, "same.png", temp_dir)
        assert exported_state[0].get("zip_report_path"), "选中项应写入单图 ZIP 路径"
        assert not exported_state[1].get("zip_report_path"), "同名但未选中的兄弟项不应被误写入 ZIP 路径"
        assert exported_state[0].get("report_path") == exported_state[0].get("zip_report_path"), "选中项报告路径应同步"
        assert not exported_state[1].get("report_path"), "未选中兄弟项 report_path 不应被污染"
    print("✓ 单图导出路径精确同步正常")
except Exception as e:
    print(f"✗ 单图导出路径同步测试失败: {e}")
    sys.exit(1)

print("\n测试42: 检查单图Word报告空模型明细回退顶层检测框...")
try:
    from docx import Document
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    detections = [{"class": "Caries", "confidence": 0.86, "x1": 1, "y1": 2, "x2": 20, "y2": 30}]
    with TemporaryDirectory() as temp_dir:
        report_path = export_single_docx_report(
            SingleReportData(
                image_name="fallback-docx.png",
                created_at="2026-07-06T10:00:00",
                model_name="primary-model",
                original_image=image,
                model_input_image=image,
                annotated_image=image,
                detections=detections,
                advice="测试建议",
                quality_text="测试质量",
                summary={},
                safety_notice="测试声明",
                model_results=[{"model": "stale-empty-model", "detections": [], "annotated": image}],
            ),
            Path(temp_dir),
        )
        doc = Document(report_path)
        doc_text = "\n".join(
            [paragraph.text for paragraph in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
        assert "primary-model" in doc_text, "空模型明细不应覆盖真实主模型名"
        assert "Caries" in doc_text and "龋齿" in doc_text, "空模型明细不应覆盖顶层检测框"
        assert "stale-empty-model" not in doc_text, "陈旧空模型名不应进入报告"
    print("✓ 单图Word报告空模型明细回退正常")
except Exception as e:
    print(f"✗ 单图Word报告空模型明细回退测试失败: {e}")
    sys.exit(1)

print("\n测试43: 检查旧模型明细缺数量时按检测框回填...")
try:
    from src.dental_detection.record_formatters import format_case_record, format_history_record

    legacy = {
        "created_at": "2026-07-06T10:00:00",
        "case_id": "legacy-count",
        "image_name": "legacy-count.png",
        "model": "legacy-model",
        "model_results": [
            {
                "model": "detail-model",
                "detections": [
                    {"class": "Impacted", "confidence": 0.77, "x1": 1, "y1": 2, "x2": 30, "y2": 40}
                ],
            }
        ],
    }
    history_detail = format_history_record(legacy)
    case_detail = format_case_record(legacy)
    assert "detail-model | 检测数量=1" in history_detail, "历史详情应按检测框数量回填模型明细数量"
    assert "detail-model | 检测数量=1" in case_detail, "病例详情应按检测框数量回填模型明细数量"
    assert "检测数量=0" not in history_detail + case_detail, "存在检测框时不应误显示数量为0"
    print("✓ 旧模型明细数量回填正常")
except Exception as e:
    print(f"✗ 旧模型明细数量回填测试失败: {e}")
    sys.exit(1)

print("\n测试44: 检查默认应用目录迁到受管子目录不递归搬动新目录...")
try:
    import src.dental_detection.assistant as assistant_module

    original_app_home = assistant_module.APP_HOME
    try:
        with TemporaryDirectory() as temp_dir:
            old_root = Path(temp_dir) / "app_home"
            old_cases = old_root / "cases"
            old_cases.mkdir(parents=True)
            (old_cases / "case_legacy.json").write_text("{}", encoding="utf-8")
            new_storage = old_cases / "nested_storage"
            assistant_module.APP_HOME = old_root

            assistant_module.migrate_storage(str(old_root), str(new_storage))

            migrated_case = new_storage / "cases" / "case_legacy.json"
            recursive_target = new_storage / "cases" / "nested_storage"
            assert migrated_case.exists(), "默认目录下的病例文件应迁入新存储目录"
            assert not recursive_target.exists(), "迁移默认目录时不应把新目录搬进自身"
    finally:
        assistant_module.APP_HOME = original_app_home
    print("✓ 默认应用目录迁到受管子目录正常")
except Exception as e:
    print(f"✗ 默认应用目录迁移到子目录测试失败: {e}")
    sys.exit(1)

print("\n测试45: 检查模型类别名兼容列表格式...")
try:
    from app import _class_name_mapping
    from src.dental_detection.inference import DentalDetector

    class FakeValue:
        def __init__(self, value):
            self._value = value

        def item(self):
            return self._value

    class FakeXYXY:
        def __getitem__(self, index):
            return self

        def tolist(self):
            return [1, 2, 30, 40]

    class FakeBox:
        cls = FakeValue(1)
        conf = FakeValue(0.82)
        xyxy = FakeXYXY()

    detector = object.__new__(DentalDetector)
    detector.names = ["Caries", "Impacted"]
    detections = detector._parse_result(type("FakeResult", (), {"boxes": [FakeBox()]})())
    assert detections[0].label == "Impacted", "names 为列表时应按类别下标读取名称"
    assert _class_name_mapping(["Caries", "Impacted"]) == {"0": "Caries", "1": "Impacted"}, (
        "导出摘要中的类别映射应兼容列表格式 names"
    )
    print("✓ 模型类别名列表格式兼容正常")
except Exception as e:
    print(f"✗ 模型类别名列表格式兼容测试失败: {e}")
    sys.exit(1)

print("\n测试46: 检查同名批量图片报告保留列表显示名...")
try:
    from docx import Document
    from PIL import Image
    import app
    from src.dental_detection.case_store import export_case_report

    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "same.png",
        "display_name": "002 - same.png",
        "result": {
            "model": "model-a",
            "detections": [],
            "original": image,
            "model_input": image,
            "annotated": image,
            "full_annotated": image,
            "table": [],
        },
        "advice": "测试建议",
        "quality_text": "测试质量",
        "quality_level": "良好",
        "summary": {},
    }
    with TemporaryDirectory() as temp_dir:
        _, _, state = app.export_word_report([item], "002 - same.png", temp_dir)
        word_path = Path(state[0]["word_report_path"])
        word_text = "\n".join(paragraph.text for paragraph in Document(word_path).paragraphs)
        assert "002 - same.png" in word_text and "原始文件：same.png" in word_text, (
            "单图 Word 报告应保留批量列表显示名，避免同名图片混淆"
        )

        ensure_app_dirs(temp_dir)
        case_path = case_dir(temp_dir) / "case_20260706_duplicate.json"
        case_path.write_text(
            json.dumps(
                {
                    "created_at": "2026-07-06T10:00:00",
                    "case_id": "duplicate",
                    "image_name": "same.png",
                    "display_name": "002 - same.png",
                    "detections": [],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        case_report = export_case_report(temp_dir, case_path.name)
        case_text = "\n".join(paragraph.text for paragraph in Document(case_report).paragraphs)
        assert "图片名称：same.png" in case_text and "列表显示名：002 - same.png" in case_text, (
            "病例 Word 报告应保留列表显示名"
        )
    print("✓ 同名批量图片报告显示名保留正常")
except Exception as e:
    print(f"✗ 同名批量图片报告显示名测试失败: {e}")
    sys.exit(1)

print("\n测试47: 检查批量导出总览保留失败和未检出信息...")
try:
    from docx import Document
    from PIL import Image
    import app
    from src.dental_detection.batch_overview_view import batch_overview_csv_text

    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "same.png",
        "display_name": "002 - same.png",
        "result": {
            "model": "model-a",
            "detections": [],
            "original": image,
            "model_input": image,
            "annotated": image,
            "full_annotated": image,
            "table": [],
        },
        "all_results": [
            {
                "model": "model-a",
                "detections": [],
                "original": image,
                "model_input": image,
                "annotated": image,
                "full_annotated": image,
            }
        ],
        "advice": "测试建议",
        "quality_text": "测试质量",
        "quality_level": "良好",
        "summary": {"模型": "model-a", "CLAHE增强": False, "conf": 0.25, "iou": 0.7},
        "batch_errors": ["bad.png: 图片处理失败"],
    }
    overview = build_batch_summary([item], item["batch_errors"])
    overview_csv = batch_overview_csv_text(overview)
    assert "无检测结果图片" in overview_csv and "002 - same.png" in overview_csv, (
        "批量总览 CSV 应包含未检出图片列表"
    )
    assert "失败图片" in overview_csv and "bad.png: 图片处理失败" in overview_csv, (
        "批量总览 CSV 应包含失败图片列表"
    )

    with TemporaryDirectory() as temp_dir:
        word_path = export_batch_docx_report([item], overview, Path(temp_dir) / "word")
        word_text = "\n".join(paragraph.text for paragraph in Document(word_path).paragraphs)
        assert "002 - same.png" in word_text and "原始文件名：same.png" in word_text, (
            "批量 Word 逐图结果应保留列表显示名和原始文件名"
        )

        _, _, state = app.export_batch_results([item], temp_dir)
        batch_zip = Path(state[0]["zip_report_path"])
        with zipfile.ZipFile(batch_zip) as archive:
            summary_text = archive.read("summary.txt").decode("utf-8")
            exported_overview_csv = archive.read("批量检测总览.csv").decode("utf-8-sig")
        assert "图片总数: 2" in summary_text and "成功处理: 1" in summary_text and "处理失败: 1" in summary_text, (
            "ZIP summary.txt 应区分总数、成功和失败数量"
        )
        assert "失败图片" in exported_overview_csv and "bad.png: 图片处理失败" in exported_overview_csv, (
            "ZIP 内批量总览 CSV 应导出失败图片列表"
        )
    print("✓ 批量导出失败和未检出信息保留正常")
except Exception as e:
    print(f"✗ 批量导出总览信息测试失败: {e}")
    sys.exit(1)

print("\n测试48: 检查移动端宽表格保留横向滚动宽度...")
try:
    css_text = load_workbench_css()
    media_anchor = "@media (max-width: 640px)"
    media_start = css_text.index(media_anchor)
    table_rule_start = css_text.index("  table {", media_start)
    table_rule_end = css_text.index("  }", table_rule_start)
    table_rule = css_text[table_rule_start:table_rule_end]
    assert "width: max-content" in table_rule, "小屏宽表格应保持内容宽度，由外层容器横向滚动"
    assert "min-width: 560px" in table_rule, "小屏检测明细表不应被压缩到100%宽"
    assert "min-width: 100%" not in table_rule, "小屏宽表格不应强制压缩到视口宽度"
    overview_rule_start = css_text.index("  .batch-overview-panel table {", table_rule_end)
    overview_rule_end = css_text.index("  }", overview_rule_start)
    overview_rule = css_text[overview_rule_start:overview_rule_end]
    assert "min-width: 720px" in overview_rule, "批量总览表列更多，应保留更宽滚动宽度"
    print("✓ 移动端宽表格横向滚动宽度正常")
except Exception as e:
    print(f"✗ 移动端宽表格样式测试失败: {e}")
    sys.exit(1)

print("\n测试49: 检查移动端顶部状态区保持紧凑...")
try:
    css_text = load_workbench_css()
    assert ".guide-steps" not in css_text, "已移除的流程提示不应继续占用首屏空间"
    media_start = css_text.index("@media (max-width: 640px)")
    status_rule_start = css_text.index("  .result-stage-note,", media_start)
    status_rule_end = css_text.index("  }", status_rule_start)
    status_rule = css_text[status_rule_start:status_rule_end]
    help_rule_start = css_text.index("  .context-help-bubble {", media_start)
    help_rule_end = css_text.index("  }", help_rule_start)
    help_rule = css_text[help_rule_start:help_rule_end]
    workbench_source = (Path(__file__).resolve().parent / "src" / "dental_detection" / "ui_workbench_page.py").read_text(encoding="utf-8")
    assert "grid-template-columns: minmax(0, 1fr) auto" in status_rule, "模型状态与提示图标应保持紧凑双列"
    assert "calc(100vw - 64px)" in help_rule, "移动端提示气泡应保留内容区安全边距"
    assert 'gr.Accordion("识别说明"' not in workbench_source, "识别说明不应退回展开式折叠面板"
    print("✓ 移动端顶部状态区收敛正常")
except Exception as e:
    print(f"✗ 移动端顶部状态区样式测试失败: {e}")
    sys.exit(1)

print("\n测试50: 检查病例无效选择在各入口提示一致...")
try:
    import app

    with TemporaryDirectory() as temp_dir:
        invalid_choice = "2026-07-06 | 病例 | 图片 | ..\\secret.json"
        detail = app.load_case_record(invalid_choice, temp_dir)
        assert "病例选择无效" in detail, "加载病例应提示选择无效"
        for action_name, action in [
            ("导出病例", lambda: app.export_selected_case_record(invalid_choice, temp_dir)),
            (
                "删除病例",
                lambda: app.delete_selected_case_record(invalid_choice, "", "全部", "全部", "", "", temp_dir),
            ),
        ]:
            try:
                action()
            except Exception as exc:
                assert "病例选择无效" in str(exc), f"{action_name} 应提示选择无效"
            else:
                raise AssertionError(f"{action_name} 不应接受无效病例选择")
    print("✓ 病例无效选择提示一致正常")
except Exception as e:
    print(f"✗ 病例无效选择提示一致测试失败: {e}")
    sys.exit(1)

print("\n测试51: 检查类别显示开关不删减检测表完整结果...")
try:
    import app
    from PIL import Image

    image = Image.new("RGB", (100, 80), "white")
    detections = [
        {"class": "Caries", "confidence": 0.88, "x1": 10, "y1": 10, "x2": 30, "y2": 30},
        {"class": "Impacted", "confidence": 0.77, "x1": 55, "y1": 35, "x2": 85, "y2": 65},
    ]
    batch_state = [
        {
            "name": "当前单图",
            "result": {
                "original": image,
                "model_input": image,
                "annotated": image,
                "detections": detections,
            },
            "advice": "",
            "summary": {},
        }
    ]
    outputs = app.update_detection_visibility(["龋齿"], "当前单图", batch_state)
    crop_items = outputs[2]
    table = outputs[5]
    assert len(crop_items) == 1, "类别开关应只影响结果图和局部图显示"
    assert set(table["class"].tolist()) == {"Caries", "Impacted"}, "检测表应保留完整检测结果"
    assert len(batch_state[0]["result"]["detections"]) == 2, "类别开关不应删除原始检测框"
    print("✓ 类别显示开关完整结果保留正常")
except Exception as e:
    print(f"✗ 类别显示开关完整结果测试失败: {e}")
    sys.exit(1)

print("\n测试52: 检查病例保存兼容仅含 image_name 的结果项...")
try:
    import app
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    item = {
        "image_name": "legacy_only_name.png",
        "result": {
            "model": "legacy-model",
            "original": image,
            "detections": [{"class": "Caries", "confidence": 0.8, "x1": 5, "y1": 6, "x2": 30, "y2": 28}],
        },
        "advice": "建议复查。",
        "summary": {},
    }
    with TemporaryDirectory() as temp_dir:
        app.save_case_record([item], "legacy_only_name.png", "", "", temp_dir)
        rows = list_case_records(temp_dir)
        assert rows and rows[0]["图片名称"] == "legacy_only_name.png", "病例列表应保留 image_name"
        data = load_case_record(temp_dir, rows[0]["文件名"])
        assert data["image_name"] == "legacy_only_name.png", "病例 JSON 不应把 image_name 回退成当前单图"
        assert data["display_name"] == "legacy_only_name.png", "病例显示名应兼容 image_name"
    print("✓ 仅 image_name 病例保存兼容正常")
except Exception as e:
    print(f"✗ 仅 image_name 病例保存兼容测试失败: {e}")
    sys.exit(1)

print("\n测试53: 检查模型格式错误提示不会泛化成加载失败...")
try:
    from src.dental_detection.error_messages import friendly_error_message

    message = friendly_error_message(
        "主模型: unsupported model format .txt; supported: .pt",
        "模型文件格式不支持",
    )
    assert "格式不受支持" in message, "模型格式错误应明确提示格式不支持"
    assert ".pt" in message and ".onnx" in message, "提示应给出可选择的模型格式"
    assert "模型文件无法加载" not in message, "格式错误不应泛化为模型损坏或无法加载"
    print("✓ 模型格式错误提示正常")
except Exception as e:
    print(f"✗ 模型格式错误提示测试失败: {e}")
    sys.exit(1)

print("\n测试54: 检查 AI 建议失败提示包含处理建议...")
try:
    import app
    from src.dental_detection.assistant import AiSettings

    original_chat_completion = app.chat_completion

    def fail_chat(*args, **kwargs):
        raise ValueError("API Key missing")

    app.chat_completion = fail_chat
    try:
        settings = AiSettings(enabled=True, api_key="", key_mode="直接 Key 值")
        advice = app._build_advice(settings, [{"class": "Caries", "confidence": 0.86}])
    finally:
        app.chat_completion = original_chat_completion
    assert "检测摘要" in advice, "AI 失败时仍应保留内置建议"
    assert "AI 接口鉴权失败" in advice, "AI 失败应转换为用户友好提示"
    assert "建议处理" in advice and "测试接口" in advice, "AI 失败提示应包含下一步操作"
    print("✓ AI 建议失败提示正常")
except Exception as e:
    print(f"✗ AI 建议失败提示测试失败: {e}")
    sys.exit(1)

print("\n测试55: 检查 AI 对话失败提示保留输入并给出建议...")
try:
    import app

    original_chat_completion = app.chat_completion

    def fail_chat(*args, **kwargs):
        raise ValueError("API Key missing")

    app.chat_completion = fail_chat
    try:
        chatbot, chat_state, chat_input, export_file, export_path = app.continue_chat(
            "请解释结果",
            [],
            True,
            "https://api.deepseek.com/v1",
            "deepseek-chat",
            "直接 Key 值",
            "",
            "",
            "",
            False,
            False,
            False,
            "",
            "",
            "简洁版",
        )
    finally:
        app.chat_completion = original_chat_completion
    assert chat_input == "请解释结果", "AI 对话失败时应保留用户输入方便修改重试"
    assert chatbot[-1]["role"] == "assistant", "失败提示应作为助手回复显示"
    assert "AI 接口鉴权失败" in chatbot[-1]["content"], "AI 对话失败应使用友好提示"
    assert "建议处理" in chatbot[-1]["content"], "AI 对话失败提示应包含处理建议"
    assert export_path == "" and getattr(export_file, "get", lambda *_: None)("visible") is False, "失败后不应暴露旧导出文件"
    print("✓ AI 对话失败提示正常")
except Exception as e:
    print(f"✗ AI 对话失败提示测试失败: {e}")
    sys.exit(1)

print("\n测试56: 检查 AI 超时和连接失败提示有明确建议...")
try:
    from src.dental_detection.error_messages import friendly_error_message

    timeout_message = friendly_error_message("AI 服务响应超时，请检查网络或接口配置。", "AI 接口测试失败")
    connect_message = friendly_error_message("无法连接 AI 服务，请检查网络、Base URL 或代理配置。", "AI 接口测试失败")
    assert "AI 服务响应超时" in timeout_message and "稍后重试" in timeout_message, (
        "AI 超时应给出重试和网络检查建议"
    )
    assert "无法连接 AI 服务" in connect_message and "测试接口" in connect_message, (
        "AI 连接失败应提示检查 Base URL/代理并测试接口"
    )
    assert "错误信息" not in timeout_message + connect_message, "常见 AI 网络错误不应落入泛化错误模板"
    print("✓ AI 网络错误提示正常")
except Exception as e:
    print(f"✗ AI 网络错误提示测试失败: {e}")
    sys.exit(1)

print("\n测试57: 检查保存设置时存储目录误填文件路径有明确提示...")
try:
    import app

    original_save_settings = app.save_settings

    def fail_save_settings(*args, **kwargs):
        raise FileExistsError("File exists: C:/tmp/not_a_dir")

    app.save_settings = fail_save_settings
    try:
        try:
            app.save_ui_settings(
                False,
                "https://api.deepseek.com/v1",
                "deepseek-chat",
                "环境变量",
                "DEEPSEEK_API_KEY",
                "",
                "",
                False,
                False,
                True,
                "C:/tmp/not_a_dir",
                "",
                "简洁版",
                False,
                False,
                app.MODEL_MODE_SINGLE,
                "models",
                str(DEFAULT_MODEL_PATH),
                "",
                True,
                100,
            )
        except Exception as exc:
            message = str(exc)
        else:
            raise AssertionError("保存设置遇到文件路径冲突时应报错")
    finally:
        app.save_settings = original_save_settings
    assert "请选择一个文件夹作为存储目录" in message, "存储目录误填文件时应明确说明选择文件夹"
    assert "建议处理" in message, "存储目录错误应包含处理建议"
    print("✓ 保存设置存储目录错误提示正常")
except Exception as e:
    print(f"✗ 保存设置存储目录错误提示测试失败: {e}")
    sys.exit(1)

print("\n测试58: 检查空字典检测框不会被统计成未知病变...")
try:
    import app
    from docx import Document
    from PIL import Image
    from src.dental_detection.history_store import build_history_record

    image = Image.new("RGB", (80, 60), "white")
    empty_detection_item = {
        "name": "empty-dict.png",
        "result": {
            "model": "model-empty",
            "original": image,
            "model_input": image,
            "annotated": image,
            "detections": [{}, {"备注": ""}],
        },
        "all_results": [{"model": "model-empty", "detections": [{}, {"备注": ""}]}],
        "summary": {},
        "advice": "",
    }

    assert app._clean_detection_records([{}, {"备注": ""}]) == [], "空字典检测框应被清理"
    summary = build_batch_summary([empty_detection_item])
    assert summary["检测框总数"] == 0 and summary["涉及类别"] == "无", "批量摘要不应把空字典当成未知病变"
    history_record = build_history_record(empty_detection_item)
    assert history_record["detection_count"] == 0, "历史记录不应统计空字典检测框"

    with TemporaryDirectory() as temp_dir:
        case_dir(temp_dir).mkdir(parents=True, exist_ok=True)
        case_path = case_dir(temp_dir) / "case_20260706_empty.json"
        case_path.write_text(
            json.dumps(
                {
                    "created_at": "2026-07-06T12:00:00",
                    "case_id": "empty",
                    "image_name": "empty-dict.png",
                    "detections": [{}],
                    "model_results": [{"model": "model-empty", "detections": [{}]}],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        rows = list_case_records(temp_dir)
        assert rows[0]["检测数量"] == 0 and rows[0]["涉及类别"] == "无检测结果", (
            "病例列表不应把空字典检测框统计为未知类别"
        )
        detail = format_case_record(load_case_record(temp_dir, case_path.name))
        assert "检测框：共 0 个" in detail and "未检测到病变框" in detail, "病例详情应显示无检测框"

        report_path = export_single_docx_report(
            SingleReportData(
                image_name="empty-dict.png",
                created_at="2026-07-06T12:00:00",
                model_name="model-empty",
                original_image=image,
                model_input_image=image,
                annotated_image=image,
                detections=[{}],
                advice="",
                quality_text="",
                summary={},
                safety_notice="测试声明",
            ),
            Path(temp_dir) / "word-empty",
        )
        report_text = "\n".join(paragraph.text for paragraph in Document(report_path).paragraphs)
        assert "未检测到明确目标框" in report_text, "Word 报告摘要不应把空字典检测框当成有效目标"
    print("✓ 空字典检测框清理正常")
except Exception as e:
    print(f"✗ 空字典检测框清理测试失败: {e}")
    sys.exit(1)

print("\n测试59: 检查仅置信度字段不会被统计成未知病变...")
try:
    import app
    from src.dental_detection.result_levels import has_detection_payload
    from src.dental_detection.history_store import build_history_record

    confidence_only = {"confidence": 0.91}
    bbox_only = {"x1": 1, "y1": 2, "x2": 20, "y2": 30}
    chinese_label = {"中文名称": "龋齿", "confidence": 0.88}
    assert not has_detection_payload(confidence_only), "只有置信度没有类别或 bbox 时不应算作检测框"
    assert has_detection_payload(bbox_only), "保留无类别但有完整 bbox 的旧记录兼容"
    assert has_detection_payload(chinese_label), "只含中文类别名的旧记录应继续兼容"
    assert app._clean_detection_records([confidence_only]) == [], "置信度孤立字段应被清理"

    item = {"name": "confidence-only.png", "result": {"detections": [confidence_only]}}
    summary = build_batch_summary([item])
    history_record = build_history_record(item)
    assert summary["检测框总数"] == 0 and history_record["detection_count"] == 0, (
        "置信度孤立字段不应污染摘要和历史统计"
    )
    print("✓ 仅置信度字段清理正常")
except Exception as e:
    print(f"✗ 仅置信度字段清理测试失败: {e}")
    sys.exit(1)

print("\n测试60: 检查模型卡片选择同步高级下拉框...")
try:
    import app

    with TemporaryDirectory() as temp_dir:
        outputs = app.apply_model_card(str(DEFAULT_MODEL_PATH), "主模型", temp_dir)
    primary_update, compare_update, dropdown_update, cards_html, info_text, feedback = outputs
    assert primary_update["value"] == str(DEFAULT_MODEL_PATH), "模型卡片应填入主模型路径"
    assert "value" not in compare_update, "选择主模型卡片不应覆盖对比模型路径"
    assert dropdown_update["value"] == str(DEFAULT_MODEL_PATH), "高级模型下拉框应同步选中模型卡片路径"
    assert any(value == str(DEFAULT_MODEL_PATH) for _, value in dropdown_update["choices"]), (
        "即使当前模型目录不包含卡片模型，也应临时补入下拉候选"
    )
    assert "selected" in cards_html and "YOLOv8m C2f-Faster-lite" in info_text, "模型信息区应同步当前卡片选择"
    assert "填入主模型" in feedback, "卡片选择应返回明确反馈"
    print("✓ 模型卡片与高级下拉同步正常")
except Exception as e:
    print(f"✗ 模型卡片与高级下拉同步测试失败: {e}")
    sys.exit(1)

print("\n测试61: 检查无效 all_results 不覆盖主结果检测框...")
try:
    from src.dental_detection.batch_summary import build_batch_summary
    from src.dental_detection.history_store import build_history_record
    from src.dental_detection.result_items import item_model_results

    primary_detection = {"class": "Caries", "confidence": 0.83, "x1": 1, "y1": 2, "x2": 30, "y2": 40}
    item = {
        "name": "primary-valid.png",
        "result": {"model": "primary-model", "detections": [primary_detection]},
        "all_results": [{"model": "stale-model", "detections": [{"confidence": 0.99}]}],
    }
    results = item_model_results(item)
    assert len(results) == 1 and results[0]["model"] == "primary-model", (
        "孤立置信度 all_results 不应覆盖有效主结果"
    )
    summary = build_batch_summary([item])
    history_record = build_history_record(item)
    assert summary["检测框总数"] == 1 and summary["涉及类别"] == "龋齿", "批量摘要应回退统计主结果"
    assert history_record["detection_count"] == 1 and history_record["model"] == "primary-model", (
        "历史记录应回退保存主结果模型和检测框"
    )
    print("✓ 无效 all_results 回退主结果正常")
except Exception as e:
    print(f"✗ 无效 all_results 回退测试失败: {e}")
    sys.exit(1)

print("\n测试62: 检查 mlpackage 目录模型可通过模型选择验证...")
try:
    import app

    with TemporaryDirectory() as temp_dir:
        package_dir = Path(temp_dir) / "exported_model.mlpackage"
        package_dir.mkdir()
        app._validate_model_artifact(package_dir, "测试 mlpackage")
        outputs = app.apply_selected_model(str(package_dir), "主模型")
    assert outputs[0]["value"].endswith("exported_model.mlpackage"), "mlpackage 目录应可填入主模型路径"
    assert "已填入主模型" in outputs[-1], "选择 mlpackage 目录模型应返回成功反馈"
    print("✓ mlpackage 目录模型选择验证正常")
except Exception as e:
    print(f"✗ mlpackage 目录模型验证测试失败: {e}")
    sys.exit(1)

print("\n测试63: 检查仅中文名称检测框不会退化为未知类别...")
try:
    import app
    from src.dental_detection.batch_summary import build_batch_summary
    from src.dental_detection.history_store import build_history_record

    chinese_only = {"中文名称": "龋齿", "置信度": "83%"}
    rows = app._clean_detection_records([chinese_only])
    assert rows[0]["class"] == "龋齿" and rows[0]["中文名称"] == "龋齿", (
        "仅含中文名称的旧检测记录应保留类别语义"
    )
    item = {"name": "chinese-only.png", "result": {"model": "legacy-model", "detections": [chinese_only]}}
    summary = build_batch_summary([item])
    history_record = build_history_record(item)
    assert summary["涉及类别"] == "龋齿" and summary["检测框总数"] == 1, "批量摘要不应显示未知类别"
    assert history_record["classes"] == ["龋齿"], "历史记录应保留中文类别名"
    print("✓ 仅中文名称检测框兼容正常")
except Exception as e:
    print(f"✗ 仅中文名称检测框兼容测试失败: {e}")
    sys.exit(1)

print("\n测试64: 检查模型和推理参数变化会清理陈旧检测结果...")
try:
    app_text = (Path(__file__).parent / "app.py").read_text(encoding="utf-8")
    assert "stale_result_controls = [primary_model_path, compare_model_path, conf, iou, device_choice, use_clahe]" in app_text, (
        "模型路径、阈值、设备和 CLAHE 变化都应注册陈旧结果清理"
    )
    assert "for control in stale_result_controls:" in app_text, "陈旧结果清理应统一绑定，避免漏掉单个控件"
    reset_helper = app_text[
        app_text.index("def chain_detection_result_reset") : app_text.index("# User-only listeners")
    ]
    assert "fn=clear_outputs_with_quality" in reset_helper, "统一重置链应清理旧检测/导出状态"
    assert "concurrency_id=RESULT_RESET_CONCURRENCY_ID" in reset_helper, "统一重置链应共享串行边界"
    assert 'trigger_mode="always_last"' in reset_helper, "统一重置链应丢弃过期的待处理请求"
    assert "cancels=inference_events" in app_text, "输入、模型或患者变化时应取消过期检测回写"
    assert app_text.count("concurrency_id=RESULT_RESET_CONCURRENCY_ID") >= 6, (
        "图片、批量文件、参数和模型应用变化都应接入统一结果重置队列"
    )
    assert "apply_model_btn.click(" in app_text and "apply_model_card_btn.click(" in app_text, "模型应用入口应存在"
    apply_model_section = app_text[app_text.index("apply_model_btn.click(") : app_text.index("test_model_btn.click(")]
    assert apply_model_section.count("chain_detection_result_reset") >= 2, "应用模型后不应保留旧检测结果"
    model_mode_section = app_text[
        app_text.index("model_mode_event = model_mode.input(") : app_text.index("test_btn.click(")
    ]
    assert model_mode_section.count("chain_detection_result_reset") >= 3, "模型模式变化后不应保留旧检测结果"
    print("✓ 模型和推理参数变化清理陈旧结果正常")
except Exception as e:
    print(f"✗ 模型和推理参数陈旧结果清理测试失败: {e}")
    sys.exit(1)

print("\n测试65: 检查无扩展名缺失模型路径优先提示不存在...")
try:
    import app

    missing_no_suffix = Path("Z:/definitely_missing_dental_model")
    try:
        app._validate_model_artifact(missing_no_suffix, "主模型")
    except Exception as exc:
        message = str(exc)
    else:
        raise AssertionError("不存在的模型路径不应通过校验")
    assert "相关文件不存在" in message or "模型文件不存在" in message, (
        "无扩展名但不存在的模型路径应优先提示文件不存在"
    )
    assert "格式不支持" not in message and "unsupported model format" not in message, (
        "缺失路径不应先被误判为格式不支持"
    )
    print("✓ 无扩展名缺失模型路径提示正常")
except Exception as e:
    print(f"✗ 无扩展名缺失模型路径提示测试失败: {e}")
    sys.exit(1)

print("\n测试66: 检查桌面端下拉菜单不被工具行裁剪...")
try:
    css_text = load_workbench_css()
    model_rule_start = css_text.index(".model-row {")
    model_rule_end = css_text.index("}", model_rule_start)
    model_rule = css_text[model_rule_start:model_rule_end]
    assert "overflow: visible" in model_rule, "模型下拉所在行不应裁剪下拉选项"

    dropdown_row_rule = ".compact-row,\n.model-row"
    assert dropdown_row_rule in css_text, "应为包含下拉框的紧凑行提供溢出可见兜底"
    dropdown_rule_start = css_text.index(dropdown_row_rule)
    dropdown_rule_end = css_text.index("}", dropdown_rule_start)
    dropdown_rule_body = css_text[dropdown_rule_start:dropdown_rule_end]
    assert "overflow: visible" in dropdown_rule_body, "病例筛选等紧凑下拉行不应裁剪选项菜单"

    safe_path_rule = ".path-row,\n.chat-input-row"
    safe_path_start = css_text.index(safe_path_rule)
    safe_path_end = css_text.index("}", safe_path_start)
    safe_path_body = css_text[safe_path_start:safe_path_end]
    assert "overflow: hidden" in safe_path_body, "路径/聊天行仍应保留控宽防溢出规则"
    print("✓ 桌面端下拉菜单裁剪防护正常")
except Exception as e:
    print(f"✗ 桌面端下拉菜单裁剪测试失败: {e}")
    sys.exit(1)

print("\n测试67: 检查中文类别名也能匹配专属内置建议...")
try:
    import src.dental_detection.assistant as assistant_module
    from src.dental_detection.assistant import default_advice
    from src.dental_detection.result_levels import get_class_display_name, normalize_class_name

    assert assistant_module._normalize_class_name is normalize_class_name, "内置建议应复用统一类别归一化函数"
    assert normalize_class_name("龋齿") == "Caries", "中文龋齿类别应归一到 Caries"
    assert normalize_class_name("根尖周病变") == "Periapical Lesion", "中文根尖周病变应归一到 Periapical Lesion"
    assert normalize_class_name("阻生牙") == "Impacted", "中文阻生牙应归一到 Impacted"
    assert get_class_display_name("龋齿") == "龋齿", "中文类别显示名应保持中文"

    advice = default_advice([{"中文名称": "龋齿", "置信度": "83%"}])
    assert "疑似龋坏相关区域" in advice, "中文类别名不应退回泛用建议"
    assert "检测到模型标记的可疑区域" not in advice, "已知中文类别应匹配专属建议"
    print("✓ 中文类别名专属建议匹配正常")
except Exception as e:
    print(f"✗ 中文类别名专属建议匹配测试失败: {e}")
    sys.exit(1)

print("\n测试68: 检查默认建议忽略无效检测字段...")
try:
    from src.dental_detection.assistant import default_advice

    advice = default_advice([{"confidence": 0.94}, {}, {"备注": ""}])
    assert "本次未检测到明确的目标病变框" in advice, "无效检测字段应按无检测处理"
    assert "未知区域" not in advice and "未知类别" not in advice, "孤立置信度不应生成未知病变建议"
    assert "0.94" not in advice, "无效置信度字段不应污染默认建议"
    print("✓ 默认建议无效检测字段清理正常")
except Exception as e:
    print(f"✗ 默认建议无效检测字段清理测试失败: {e}")
    sys.exit(1)

print("\n测试69: 检查 CSV 导出转义公式型文本...")
try:
    import app
    from PIL import Image
    from src.dental_detection.batch_overview_view import batch_overview_csv_text
    from src.dental_detection.text_utils import csv_safe_value

    assert csv_safe_value("=HYPERLINK(\"http://bad\")").startswith("'="), "公式型 CSV 单元格应加前缀"

    image = Image.new("RGB", (80, 60), "white")
    malicious_name = '=HYPERLINK("http://bad","x").png'
    item = {
        "name": malicious_name,
        "display_name": malicious_name,
        "result": {
            "model": "model-safe",
            "original": image,
            "model_input": image,
            "annotated": image,
            "detections": [{"class": "Caries", "confidence": 0.86, "x1": 1, "y1": 2, "x2": 30, "y2": 40}],
        },
        "summary": {"模型": "model-safe", "CLAHE增强": False, "conf": 0.25, "iou": 0.7},
        "advice": "",
    }
    with TemporaryDirectory() as temp_dir:
        _, _, state = app.export_batch_results([item], temp_dir)
        zip_path = Path(state[0]["zip_report_path"])
        with zipfile.ZipFile(zip_path) as archive:
            detections_csv = archive.read("detections.csv").decode("utf-8-sig")
            overview_csv = archive.read("批量检测总览.csv").decode("utf-8-sig")
    assert "'=HYPERLINK" in detections_csv, "检测明细 CSV 中的图片名应防公式注入"
    assert "'=HYPERLINK" in overview_csv, "批量总览 CSV 中的显示名应防公式注入"
    direct_overview = batch_overview_csv_text({"失败图片": [malicious_name]})
    assert "'=HYPERLINK" in direct_overview, "直接生成总览 CSV 时也应转义公式型文本"
    print("✓ CSV 公式型文本转义正常")
except Exception as e:
    print(f"✗ CSV 公式型文本转义测试失败: {e}")
    sys.exit(1)

print("\n测试70: 检查单图 ZIP 导出兼容异常摘要结构...")
try:
    import app
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "summary-list.png",
        "result": {
            "original": image,
            "model_input": image,
            "annotated": image,
            "detections": [],
        },
        "summary": ["bad legacy summary"],
        "advice": "",
    }
    with TemporaryDirectory() as temp_dir:
        _, _, state = app.export_single_report([item], "summary-list.png", temp_dir)
        zip_path = Path(state[0]["zip_report_path"])
        with zipfile.ZipFile(zip_path) as archive:
            payload = json.loads(archive.read("detections.json").decode("utf-8"))

        malformed_model_item = {
            **item,
            "name": "summary-model-list.png",
            "summary": {"模型结果": ["bad legacy summary"]},
        }
        _, _, malformed_state = app.export_single_report([malformed_model_item], "summary-model-list.png", temp_dir)
        malformed_zip_path = Path(malformed_state[0]["zip_report_path"])
        with zipfile.ZipFile(malformed_zip_path) as archive:
            malformed_payload = json.loads(archive.read("detections.json").decode("utf-8"))
        _, _, word_state = app.export_word_report([malformed_model_item], "summary-model-list.png", temp_dir)
        word_path = Path(word_state[0]["word_report_path"])
        word_exists = word_path.exists()
    assert payload["summary"] == {}, "异常 summary 结构应被归一为空摘要，避免导出崩溃"
    assert payload["report"]["model"] == "unknown", "缺失模型名时应回退 unknown"
    assert malformed_payload["report"]["model"] == "unknown", "异常模型明细不应导致模型名回填崩溃"
    assert word_exists, "Word 报告也应兼容异常模型明细"
    print("✓ 单图 ZIP 异常摘要结构兼容正常")
except Exception as e:
    print(f"✗ 单图 ZIP 异常摘要结构测试失败: {e}")
    sys.exit(1)

print("\n测试71: 检查批量 Word 报告跳过异常模型摘要项...")
try:
    from docx import Document
    from PIL import Image
    from src.dental_detection.reporting import export_batch_docx_report

    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "batch-summary.png",
        "result": {
            "original": image,
            "model_input": image,
            "annotated": image,
            "detections": [],
        },
        "summary": {"模型结果": ["bad legacy summary", {"模型": "valid-model"}]},
        "advice": "",
    }
    with TemporaryDirectory() as temp_dir:
        path = export_batch_docx_report([item], None, Path(temp_dir) / "word")
        paragraphs = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    assert "使用模型：valid-model" in paragraphs, "批量 Word 报告应跳过异常摘要项并使用后续有效模型名"
    print("✓ 批量 Word 异常模型摘要兼容正常")
except Exception as e:
    print(f"✗ 批量 Word 异常模型摘要测试失败: {e}")
    sys.exit(1)

print("\n测试72: 检查模型信息不把无效模型路径标为可用...")
try:
    import app

    with TemporaryDirectory() as temp_dir:
        fake_pt_dir = Path(temp_dir) / "fake_model.pt"
        fake_pt_dir.mkdir()
        fake_package_file = Path(temp_dir) / "fake_package.mlpackage"
        fake_package_file.write_bytes(b"not a package directory")

        pt_info = app._current_model_info_markdown(str(fake_pt_dir))
        package_info = app._current_model_info_markdown(str(fake_package_file))
    assert "模型状态：不可用或格式不支持" in pt_info, ".pt 目录不应在模型信息中显示为可用"
    assert "模型状态：不可用或格式不支持" in package_info, ".mlpackage 普通文件不应显示为可用模型包"
    assert "模型状态：可用" not in pt_info + package_info, "无效模型路径不能误导用户为可用"
    print("✓ 模型信息无效路径状态正常")
except Exception as e:
    print(f"✗ 模型信息无效路径状态测试失败: {e}")
    sys.exit(1)

print("\n测试73: 检查病例详情兼容异常摘要结构...")
try:
    import app
    from src.dental_detection.assistant import case_dir, ensure_app_dirs
    from src.dental_detection.record_formatters import format_case_record

    malformed_case = {
        "case_id": "legacy-bad-summary",
        "created_at": "2026-07-06T12:45:00",
        "image_name": "bad-summary.png",
        "summary": ["bad legacy summary"],
        "detections": [],
    }
    detail = format_case_record(malformed_case)
    assert "病例编号：legacy-bad-summary" in detail, "异常 summary 不应阻断病例详情基础信息展示"
    assert "暂无摘要信息" in detail, "异常 summary 应回退为空摘要提示"

    with TemporaryDirectory() as temp_dir:
        ensure_app_dirs(temp_dir)
        case_path = case_dir(temp_dir) / "case_bad_summary.json"
        case_path.write_text(json.dumps(malformed_case, ensure_ascii=False), encoding="utf-8")
        ui_detail = app.load_case_record(
            "2026-07-06T12:45:00 | legacy-bad-summary | bad-summary.png | case_bad_summary.json",
            temp_dir,
        )
    assert "错误" not in ui_detail and "legacy-bad-summary" in ui_detail, (
        "病例详情入口应兼容有效 JSON 中的异常 summary 结构"
    )
    print("✓ 病例详情异常摘要结构兼容正常")
except Exception as e:
    print(f"✗ 病例详情异常摘要结构测试失败: {e}")
    sys.exit(1)

print("\n测试74: 检查非对象病例 JSON 显示为损坏记录...")
try:
    import app
    from src.dental_detection.assistant import case_dir, ensure_app_dirs
    from src.dental_detection.case_store import list_case_records as list_case_rows

    with TemporaryDirectory() as temp_dir:
        ensure_app_dirs(temp_dir)
        case_path = case_dir(temp_dir) / "case_list_payload.json"
        case_path.write_text("[1, 2, 3]", encoding="utf-8")
        rows = list_case_rows(temp_dir)
        detail = app.load_case_record("x | y | z | case_list_payload.json", temp_dir)
        try:
            app.export_selected_case_record("x | y | z | case_list_payload.json", temp_dir)
        except Exception as exc:
            export_message = str(exc)
        else:
            raise AssertionError("非对象病例 JSON 不应导出为空报告")
    assert rows and rows[0]["病例编号"] == "损坏病例文件", "顶层非对象病例 JSON 应显示为损坏记录"
    assert "病例文件损坏或无法读取" in detail, "详情入口应提示病例文件损坏而不是显示空病例"
    assert "病例文件损坏或无法读取" in export_message, "导出入口应阻止非对象病例 JSON 生成空报告"
    print("✓ 非对象病例 JSON 损坏提示正常")
except Exception as e:
    print(f"✗ 非对象病例 JSON 损坏提示测试失败: {e}")
    sys.exit(1)

print("\n测试75: 检查单对象检测框兼容为一条结果...")
try:
    import app
    from PIL import Image
    from src.dental_detection.assistant import default_advice, detection_prompt
    from src.dental_detection.batch_summary import build_batch_summary
    from src.dental_detection.history_store import build_history_record
    from src.dental_detection.record_formatters import format_case_record, format_history_record
    from src.dental_detection.visualization import crop_detection_regions

    detection = {"class": "Caries", "confidence": 0.86, "x1": 4, "y1": 5, "x2": 40, "y2": 45}
    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "single-dict-detection.png",
        "result": {"model": "legacy-model", "detections": detection, "original": image},
        "all_results": [{"model": "legacy-model", "detections": detection}],
    }

    assert len(app._clean_detection_records(detection)) == 1, "单个检测框对象应被当作一条检测结果"
    assert "龋齿" in default_advice(detection), "内置建议应兼容单对象检测框"
    assert '"class": "Caries"' in detection_prompt(detection)[1]["content"], "AI 提示应把单对象检测框转为列表内容"
    assert build_batch_summary([item])["检测框总数"] == 1, "批量摘要应统计单对象检测框"

    history_record = build_history_record(item)
    assert history_record["detection_count"] == 1, "历史记录应统计单对象检测框"
    assert "龋齿" in format_history_record(history_record), "历史详情应显示单对象检测框"

    case_detail = format_case_record({"case_id": "single-dict", "image_name": "x.png", "model_results": item["all_results"]})
    assert "检测框：共 1 个" in case_detail and "龋齿" in case_detail, "病例详情应显示单对象检测框"
    assert len(crop_detection_regions(image, detection)) == 1, "局部区域裁剪应兼容单对象检测框"
    print("✓ 单对象检测框兼容正常")
except Exception as e:
    print(f"✗ 单对象检测框兼容测试失败: {e}")
    sys.exit(1)

print("\n测试76: 检查单对象模型结果兼容为一组结果...")
try:
    import app
    from docx import Document
    from PIL import Image
    from src.dental_detection.case_store import _case_detections
    from src.dental_detection.history_store import build_history_record
    from src.dental_detection.record_formatters import format_case_record, format_history_record
    from src.dental_detection.reporting import export_batch_docx_report

    detection = {"class": "Caries", "confidence": 0.88, "x1": 4, "y1": 5, "x2": 40, "y2": 45}
    model_result = {"model": "single-model-object", "模型": "single-model-object", "detections": detection}
    image = Image.new("RGB", (80, 60), "white")
    item = {
        "name": "single-model-result.png",
        "result": {"model": "primary", "detections": [], "original": image, "model_input": image, "annotated": image},
        "all_results": model_result,
        "summary": {"模型结果": model_result},
        "advice": "",
    }

    assert len(app._item_results(item)) == 1 and app._item_results(item)[0]["model"] == "single-model-object", (
        "all_results 单对象应被当作一组模型结果"
    )
    history_record = build_history_record(item)
    assert history_record["detection_count"] == 1 and history_record["model"] == "single-model-object", (
        "历史记录应统计 all_results 单对象模型结果"
    )
    assert "龋齿" in format_history_record({"image_name": "x", "model_results": model_result}), (
        "历史详情应兼容 model_results 单对象"
    )
    assert len(_case_detections({"model_results": model_result})) == 1, "病例列表统计应兼容 model_results 单对象"
    assert "检测框：共 1 个" in format_case_record({"case_id": "single-model", "image_name": "x", "model_results": model_result}), (
        "病例详情应兼容 model_results 单对象"
    )
    assert app._summary_model_name({"模型结果": model_result}) == "single-model-object", (
        "摘要模型名应兼容单对象模型结果"
    )
    with TemporaryDirectory() as temp_dir:
        report_path = export_batch_docx_report([item], None, Path(temp_dir) / "word")
        paragraphs = "\n".join(paragraph.text for paragraph in Document(report_path).paragraphs)
    assert "使用模型：single-model-object" in paragraphs, "批量 Word 报告应显示单对象模型结果名称"
    print("✓ 单对象模型结果兼容正常")
except Exception as e:
    print(f"✗ 单对象模型结果兼容测试失败: {e}")
    sys.exit(1)

print("\n测试77: 检查中文字段模型结果贯穿前端导出入口...")
try:
    import app
    from docx import Document
    from PIL import Image

    image = Image.new("RGB", (80, 60), "white")
    detection = {"类别": "Caries", "置信度": "88%", "x1": 4, "y1": 5, "x2": 40, "y2": 45}
    model_result = {
        "模型": "中文字段模型",
        "检测框": detection,
        "路径": "server_models/chinese.pt",
        "annotated": image,
    }
    item = {
        "name": "chinese-model-result.png",
        "display_name": "中文字段结果",
        "result": {"模型": "主模型中文", "检测框": [], "original": image, "model_input": image, "annotated": image},
        "all_results": [model_result],
        "summary": {"模型结果": model_result, "模型": "摘要中文模型"},
        "advice": "测试建议",
        "suggestion_type": "default",
    }

    assert app._model_result_records(item)[0]["detection_count"] == 1, "病例保存入口应统计中文字段检测框"
    assert app._advice_detections([model_result])[0]["模型"] == "中文字段模型", "建议入口应保留中文字段模型名"
    assert "中文字段模型" in app._model_detections_html([model_result]) or "龋齿" in app._model_detections_html([model_result]), (
        "HTML 检测框入口应兼容中文字段模型结果"
    )
    assert app._summary_lines([item], {"exported_at": "now", "use_clahe": False, "conf": 0.25, "iou": 0.45})[
        5
    ] == "检测到的总框数: 1", "批量文本摘要应统计中文字段检测框"

    with TemporaryDirectory() as temp_dir:
        _, _, batch_state = app.export_batch_results([item], temp_dir)
        batch_zip = Path(batch_state[0]["zip_report_path"])
        with zipfile.ZipFile(batch_zip) as archive:
            batch_payload = json.loads(archive.read("detections.json").decode("utf-8"))
            csv_text = archive.read("detections.csv").decode("utf-8-sig")
        batch_model = batch_payload["items"][0]["models"][0]
        assert batch_model["model"] == "中文字段模型", "批量 ZIP 应导出中文字段模型名"
        assert batch_model["model_path"] == "server_models/chinese.pt", "批量 ZIP 应导出中文字段路径"
        assert len(batch_model["detections"]) == 1, "批量 ZIP 应导出中文字段检测框"
        assert "中文字段模型" in csv_text and "龋齿" in csv_text, "批量 CSV 应包含中文字段模型结果"

        _, _, single_state = app.export_single_report([item], "chinese-model-result.png", temp_dir)
        single_zip = Path(single_state[0]["zip_report_path"])
        with zipfile.ZipFile(single_zip) as archive:
            single_payload = json.loads(archive.read("detections.json").decode("utf-8"))
        assert single_payload["models"][0]["model"] == "中文字段模型", "单图 ZIP 应导出中文字段模型名"
        assert len(single_payload["models"][0]["detections"]) == 1, "单图 ZIP 应导出中文字段检测框"

        _, _, word_state = app.export_word_report([item], "chinese-model-result.png", temp_dir)
        paragraphs = "\n".join(paragraph.text for paragraph in Document(word_state[0]["word_report_path"]).paragraphs)
        assert "使用模型：中文字段模型" in paragraphs, "单图 Word 报告应显示中文字段模型名"
    print("✓ 中文字段模型结果前端导出入口兼容正常")
except Exception as e:
    print(f"✗ 中文字段模型结果前端导出入口测试失败: {e}")
    sys.exit(1)

print("\n测试78: 检查主流程输出不会把按钮文字传给文件组件...")
try:
    import app
    from PIL import Image

    def assert_common_outputs_aligned(outputs):
        assert len(outputs) == len(COMMON_OUTPUT_KEYS), (
            f"主流程应返回 {len(COMMON_OUTPUT_KEYS)} 个输出，实际 {len(outputs)}"
        )
        assert getattr(common_output(outputs, "word_export_button"), "get", lambda *_: None)("value") == (
            "导出 Word 报告"
        ), "Word 导出按钮应匹配命名输出契约"
        assert getattr(common_output(outputs, "zip_export_button"), "get", lambda *_: None)("value") == (
            "导出 ZIP 数据包"
        ), "ZIP 导出按钮应匹配命名输出契约"
        assert getattr(common_output(outputs, "word_report_file"), "get", lambda *_: None)("value") != (
            "导出 Word 报告"
        ), "word_report_file 文件组件不能收到按钮文字"
        assert getattr(common_output(outputs, "zip_report_file"), "get", lambda *_: None)("value") != (
            "导出 ZIP 数据包"
        ), "zip_report_file 文件组件不能收到按钮文字"

    image = Image.new("RGB", (80, 60), "white")
    detection = {
        "class": "Caries",
        "confidence": 0.86,
        "x1": 4,
        "y1": 5,
        "x2": 40,
        "y2": 45,
    }

    def fake_detect_model_path(model_name, model_path, image_arg, use_clahe, conf, iou, device):
        records = app._clean_detection_records([detection], image_size=image.size)
        return {
            "model": model_name,
            "original": image,
            "model_input": image,
            "annotated": image,
            "full_annotated": image,
            "detections": records,
            "table": app._table_from_records(records),
            "class_names": {"0": "Caries"},
            "model_path": str(model_path),
        }

    original_detect = app._detect_model_path
    app._detect_model_path = fake_detect_model_path
    try:
        baseline = str(Path("models/final_candidates/yolov8m_1280_full/weights/best.pt").resolve())
        single_outputs = app.run_single_detection(
            image,
            app.MODEL_MODE_SINGLE,
            "personal-self",
            baseline,
            baseline,
            0.25,
            0.45,
            "cpu",
            False,
            False,
            True,
            False,
            "https://api.deepseek.com/v1",
            "deepseek-chat",
            "环境变量",
            "DEEPSEEK_API_KEY",
            "",
            "",
            False,
            False,
            False,
            str(Path("outputs/test-storage").resolve()),
            "",
            "简洁版",
            False,
            100,
        )
        assert_common_outputs_aligned(single_outputs)
        single_state = common_output(single_outputs, "batch_state")
        assert single_state[0].get("task_id"), "单图检测应生成工作区任务记录"
        assert single_state[0].get("patient_id") == "personal-self", "单图任务应关联当前患者档案"

        with TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "batch.png"
            image.save(image_path)
            batch_outputs = app.run_batch_detection(
                [str(image_path)],
                app.MODEL_MODE_SINGLE,
                "personal-self",
                baseline,
                baseline,
                0.25,
                0.45,
                "cpu",
                False,
                False,
                True,
                False,
                "https://api.deepseek.com/v1",
                "deepseek-chat",
                "环境变量",
                "DEEPSEEK_API_KEY",
                "",
                "",
                False,
                False,
                False,
                str(Path(temp_dir) / "storage"),
                "",
                "简洁版",
                False,
                100,
            )
        assert_common_outputs_aligned(batch_outputs)
        batch_state = common_output(batch_outputs, "batch_state")
        assert batch_state[0].get("task_id"), "批量检测应逐图生成工作区任务记录"
        assert batch_state[0].get("patient_id") == "personal-self", "批量任务应关联当前患者档案"
        assert_common_outputs_aligned(app.clear_outputs())
    finally:
        app._detect_model_path = original_detect
    print("✓ 主流程输出组件顺序正常")
except Exception as e:
    print(f"✗ 主流程输出组件顺序测试失败: {e}")
    sys.exit(1)

print("\n测试79: 检查模型 UI helper 输出关键展示状态...")
try:
    from src.dental_detection.config import DEFAULT_MODEL_PATH, MODEL_REGISTRY, PROJECT_ROOT
    from src.dental_detection.model_info import build_model_cards
    from src.dental_detection.model_ui import (
        build_advanced_model_warning_html,
        build_model_path_compact_html,
        build_workbench_model_status_html,
        format_model_path_for_display,
    )

    cards = build_model_cards(MODEL_REGISTRY)
    recommended_paths = {str(info["path"]) for info in MODEL_REGISTRY.values()}
    baseline_path = str(MODEL_REGISTRY["YOLOv8m 原始结构"]["path"])
    optimized_path = str(MODEL_REGISTRY["YOLOv8m C2f-Faster-lite"]["path"])

    baseline_html = build_workbench_model_status_html(
        baseline_path,
        cards,
        default_model_path=str(DEFAULT_MODEL_PATH),
        recommended_paths=recommended_paths,
    )
    assert "兼容模型" in baseline_html, "原始结构状态应保留兼容模型标识"

    optimized_html = build_workbench_model_status_html(
        optimized_path,
        cards,
        default_model_path=str(DEFAULT_MODEL_PATH),
        recommended_paths=recommended_paths,
    )
    assert "优化模型" in optimized_html, "C2f 状态应保留优化模型标识"

    advanced_path = str(PROJECT_ROOT / "models" / "pretrained" / "yolov8n.pt")
    advanced_html = build_workbench_model_status_html(
        advanced_path,
        cards,
        default_model_path=str(DEFAULT_MODEL_PATH),
        recommended_paths=recommended_paths,
    )
    assert "请确认来源和兼容性" in advanced_html, "高级模型状态应提示兼容性风险"

    long_path = "E:/" + "/".join(["very_long_model_directory"] * 8) + "/weights/best.pt"
    compact_text = format_model_path_for_display(long_path, max_chars=60)
    compact_html = build_model_path_compact_html(long_path)
    assert "..." in compact_text and compact_text != long_path, "长路径展示文本应省略"
    assert f'title="{long_path}"' in compact_html, "长路径 HTML 应保留完整 title 便于追踪"
    assert build_advanced_model_warning_html().strip(), "高级模型提示不能为空"
    print("✓ 模型 UI helper 输出正常")
except Exception as e:
    print(f"✗ 模型 UI helper 测试失败: {e}")
    sys.exit(1)

print("\n测试80: 检查导出 helper 文件组织和安全写入...")
try:
    from src.dental_detection.exporters import (
        build_export_manifest,
        create_zip_from_manifest,
        safe_export_stem,
        write_csv_file,
        write_html_file,
        write_json_file,
    )

    assert safe_export_stem("") == "image", "空文件名应回退为 image"
    assert safe_export_stem("病例:测试?.png") == "病例_测试", "中文文件名应保留并过滤 Windows 禁用字符"
    assert safe_export_stem("CON.png") == "CON_file", "Windows 保留名应追加后缀"

    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        csv_path = write_csv_file(
            root / "detections.csv",
            ["name", "value"],
            [{"name": "=HYPERLINK(\"http://bad\")", "value": "+SUM(1,2)"}],
        )
        csv_text = csv_path.read_text(encoding="utf-8-sig")
        assert "'=HYPERLINK" in csv_text and "'+SUM" in csv_text, "CSV 写入应保留公式注入防护"

        json_path = write_json_file(root / "数据.json", {"中文": "正常", "列表": [1, 2]})
        assert '"中文": "正常"' in json_path.read_text(encoding="utf-8"), "JSON 写入应保留中文"

        html_path = write_html_file(root / "report.html", "<html><body>牙齿报告</body></html>")
        assert "牙齿报告" in html_path.read_text(encoding="utf-8"), "HTML 应使用 UTF-8 写入"

        manifest = build_export_manifest(root)
        archive_names = {item["archive_name"] for item in manifest}
        assert {"detections.csv", "数据.json", "report.html"}.issubset(archive_names), "manifest 应包含预期文件"

        missing = root / "missing.txt"
        zip_path, skipped = create_zip_from_manifest(
            root / "bundle.zip",
            [*manifest, {"source_path": str(missing), "archive_name": "missing.txt"}],
        )
        assert str(missing) in skipped, "缺失文件加入 ZIP 时应被记录为跳过"
        with zipfile.ZipFile(zip_path) as archive:
            names = set(archive.namelist())
        assert "detections.csv" in names and "missing.txt" not in names, "ZIP 应包含现有文件并跳过缺失文件"
    print("✓ 导出 helper 文件组织正常")
except Exception as e:
    print(f"✗ 导出 helper 测试失败: {e}")
    sys.exit(1)

print("\n" + "="*60)
print("测试完成！")
